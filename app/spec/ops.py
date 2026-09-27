"""对规格的限定操作。对话修改和直接编辑都只能通过这些操作改动规格，每个操作指明目标 ID。

应用前检查修改范围：选中范围之外的目标会被拒绝，保证“只改指定目标”。
"""
from __future__ import annotations

import copy
from typing import Any

from ..util import UserError, new_id
from .deck import Deck, Slide
from .document import Document
from .workbook import Sheet, Workbook


class OpError(UserError):
    pass


# ---------- 路径读写（content.cards.1.body 形式） ----------

def get_path(obj: Any, path: str) -> Any:
    cur = obj
    for part in path.split("."):
        if isinstance(cur, list):
            cur = cur[int(part)]
        else:
            cur = cur[part]
    return cur


def set_path(obj: Any, path: str, value: Any) -> None:
    parts = path.split(".")
    cur = obj
    for part in parts[:-1]:
        cur = cur[int(part)] if isinstance(cur, list) else cur[part]
    last = parts[-1]
    if isinstance(cur, list):
        cur[int(last)] = value
    else:
        if last not in cur and not isinstance(cur, dict):
            raise OpError(f"路径不存在：{path}")
        cur[last] = value


# ---------- PPT ----------

def apply_deck_ops(deck: Deck, ops: list[dict], scope: dict | None = None) -> tuple[Deck, list[str]]:
    """返回 (新规格, 变化的页面 ID)。scope: {type: all|slides|elements, ids: [...]}。"""
    scope = scope or {"type": "all"}
    d = deck.model_dump(mode="json")
    slides = d["slides"]
    ids = [s["id"] for s in slides]
    stype = scope.get("type", "all")
    allowed = set(scope.get("ids") or [])
    # elements 范围：ids 为 "页面ID|路径"
    elem_scope: dict[str, set[str]] = {}
    if stype == "elements":
        for x in allowed:
            sid, _, path = x.partition("|")
            elem_scope.setdefault(sid, set()).add(path)
    changed: list[str] = []

    def idx(sid):
        try:
            return ids.index(sid)
        except ValueError:
            raise OpError(f"页面不存在：{sid}")

    def need_slide(sid):
        if stype == "slides" and sid not in allowed:
            raise OpError("修改超出了选中的页面范围")
        if stype == "elements" and sid not in elem_scope:
            raise OpError("修改超出了选中的元素范围")

    for op in ops:
        kind = op.get("op")
        if kind == "set_field":
            sid, path = op["slide_id"], op["path"]
            need_slide(sid)
            if stype == "elements" and not any(path == p or path.startswith(p + ".") for p in elem_scope[sid]):
                raise OpError(f"修改超出了选中的元素范围：{path}")
            s = slides[idx(sid)]
            if path == "title" or path == "notes":
                s[path] = op["value"]
            else:
                p = path if path.startswith("content.") else "content." + path
                set_path(s, p, op["value"])
            changed.append(sid)
        elif kind == "replace_slide":
            sid = op["slide_id"]
            need_slide(sid)
            if stype == "elements":
                raise OpError("选中元素时只能修改元素内容")
            new = dict(op["slide"])
            new["id"] = sid
            slides[idx(sid)] = new
            changed.append(sid)
        elif kind == "insert_slide":
            if stype == "elements":
                raise OpError("选中元素时不能增加页面")
            after = op.get("after")
            if stype == "slides" and after not in allowed and after is not None:
                raise OpError("只能在选中页面之后插入")
            new = dict(op["slide"])
            new["id"] = new_id("s")
            pos = (idx(after) + 1) if after else 0
            slides.insert(pos, new)
            ids.insert(pos, new["id"])
            changed.append(new["id"])
        elif kind == "delete_slide":
            sid = op["slide_id"]
            need_slide(sid)
            if stype == "elements":
                raise OpError("选中元素时不能删除页面")
            i = idx(sid)
            slides.pop(i)
            ids.pop(i)
        elif kind == "move_slide":
            sid = op["slide_id"]
            need_slide(sid)
            if stype == "elements":
                raise OpError("选中元素时不能移动页面")
            i = idx(sid)
            s = slides.pop(i)
            ids.pop(i)
            to = max(0, min(len(slides), int(op["to"])))
            slides.insert(to, s)
            ids.insert(to, sid)
            changed.append(sid)
        elif kind == "set_theme":
            if stype != "all" and not scope.get("allow_global"):
                raise OpError("修改配色和字体需要选择“整份文档”")
            d["theme"] = {**d["theme"], **(op.get("theme") or {})}
            changed.extend(ids)
        elif kind == "set_deck":
            if stype != "all" and not scope.get("allow_global"):
                raise OpError("修改整份文档的设置需要选择“整份文档”")
            for k in ("title", "footer", "aspect"):
                if k in op:
                    d[k] = op[k]
            changed.extend(ids)
        else:
            raise OpError(f"不支持的操作：{kind}")
    if not slides:
        raise OpError("至少保留一页")
    new = Deck.model_validate(d)
    return new, list(dict.fromkeys(c for c in changed if c in [s.id for s in new.slides]))


# ---------- Word ----------

def apply_doc_ops(doc: Document, ops: list[dict], scope: dict | None = None) -> tuple[Document, list[str]]:
    scope = scope or {"type": "all"}
    d = doc.model_dump(mode="json")
    blocks = d["blocks"]
    ids = [b["id"] for b in blocks]
    stype = scope.get("type", "all")
    allowed = set(scope.get("ids") or [])
    changed: list[str] = []

    def idx(bid):
        try:
            return ids.index(bid)
        except ValueError:
            raise OpError(f"内容块不存在：{bid}")

    def need(bid):
        if stype == "blocks" and bid not in allowed:
            raise OpError("修改超出了选中的内容范围")

    for op in ops:
        kind = op.get("op")
        if kind == "replace_block":
            bid = op["block_id"]
            need(bid)
            nb = dict(op["block"])
            nb["id"] = bid
            blocks[idx(bid)] = nb
            changed.append(bid)
        elif kind == "set_text":
            bid = op["block_id"]
            need(bid)
            b = blocks[idx(bid)]
            if "text" not in b:
                raise OpError("该内容块没有文字字段")
            b["text"] = op["text"]
            changed.append(bid)
        elif kind == "insert_blocks":
            after = op.get("after")
            if stype == "blocks" and after is not None and after not in allowed:
                raise OpError("只能在选中内容之后插入")
            pos = (idx(after) + 1) if after else 0
            for nb in op.get("blocks") or []:
                nb = dict(nb)
                nb["id"] = new_id("b")
                blocks.insert(pos, nb)
                ids.insert(pos, nb["id"])
                changed.append(nb["id"])
                pos += 1
        elif kind == "delete_block":
            bid = op["block_id"]
            need(bid)
            i = idx(bid)
            blocks.pop(i)
            ids.pop(i)
        elif kind == "move_block":
            bid = op["block_id"]
            need(bid)
            i = idx(bid)
            b = blocks.pop(i)
            ids.pop(i)
            after = op.get("after")
            pos = (idx(after) + 1) if after else 0
            blocks.insert(pos, b)
            ids.insert(pos, bid)
            changed.append(bid)
        elif kind == "set_meta":
            if stype != "all" and not scope.get("allow_global"):
                raise OpError("修改文档设置需要选择“整份文档”")
            for k in ("title", "preset", "header_text", "font_mode", "page_numbers"):
                if k in op:
                    d[k] = op[k]
            if op.get("meta"):
                d["meta"] = {**d["meta"], **op["meta"]}
            changed.append("_meta")
        else:
            raise OpError(f"不支持的操作：{kind}")
    if not blocks:
        raise OpError("文档至少保留一个内容块")
    new = Document.model_validate(d)
    return new, list(dict.fromkeys(changed))


# ---------- Excel ----------

def _addr(cell: str) -> tuple[int, int]:
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    col, row = coordinate_from_string(cell.upper())
    return row, column_index_from_string(col)


def _in_range(sheet: str, r: int, c: int, rng: dict | None) -> bool:
    if not rng:
        return True
    return sheet == rng["sheet"] and rng["r1"] <= r <= rng["r2"] and rng["c1"] <= c <= rng["c2"]


def parse_range(s: str) -> dict:
    """'销售数据!B2:D20' → {sheet, r1, c1, r2, c2}（行列从 1 开始，含表头行）。"""
    from openpyxl.utils import range_boundaries
    sheet, _, rng = s.rpartition("!")
    c1, r1, c2, r2 = range_boundaries(rng.replace("$", ""))
    return {"sheet": sheet.strip("'"), "r1": r1, "c1": c1, "r2": r2 or r1, "c2": c2 or c1}


def apply_wb_ops(wb: Workbook, ops: list[dict], scope: dict | None = None) -> tuple[Workbook, list[str]]:
    scope = scope or {"type": "all"}
    d = wb.model_dump(mode="json")
    rng = parse_range(scope["range"]) if scope.get("type") == "range" else None
    changed: list[str] = []

    def sheet(name):
        for s in d["sheets"]:
            if s["name"] == name or s["id"] == name:
                return s
        raise OpError(f"工作表不存在：{name}")

    def put(s, r, c, v):
        if rng and not _in_range(s["name"], r, c, rng):
            raise OpError("修改超出了选中的单元格区域")
        hdr = 1 if s["columns"] else 0
        if r == 1 and hdr:
            while len(s["columns"]) < c:
                s["columns"].append({"header": ""})
            s["columns"][c - 1]["header"] = "" if v is None else str(v)
            return
        ri = r - 1 - hdr
        while len(s["rows"]) <= ri:
            s["rows"].append([])
        row = s["rows"][ri]
        while len(row) < c:
            row.append(None)
        row[c - 1] = v

    for op in ops:
        kind = op.get("op")
        if kind == "set_cell":
            s = sheet(op["sheet"])
            r, c = _addr(op["cell"])
            put(s, r, c, op.get("value"))
            changed.append(s["id"])
        elif kind == "set_range":
            s = sheet(op["sheet"])
            r0, c0 = _addr(op["start"])
            for i, rowv in enumerate(op.get("values") or []):
                for j, v in enumerate(rowv):
                    put(s, r0 + i, c0 + j, v)
            changed.append(s["id"])
        elif kind in ("replace_sheet", "add_sheet", "delete_sheet", "set_columns", "add_chart", "delete_rows", "insert_rows"):
            if rng:
                raise OpError("选中单元格区域时只能修改区域内的单元格")
            if kind == "replace_sheet":
                s = sheet(op["sheet"])
                new = dict(op["data"])
                new["id"] = s["id"]
                d["sheets"][d["sheets"].index(s)] = new
                changed.append(s["id"])
            elif kind == "add_sheet":
                new = dict(op["data"])
                new["id"] = new_id("sh")
                d["sheets"].append(new)
                changed.append(new["id"])
            elif kind == "delete_sheet":
                s = sheet(op["sheet"])
                d["sheets"].remove(s)
            elif kind == "set_columns":
                s = sheet(op["sheet"])
                s["columns"] = op["columns"]
                changed.append(s["id"])
            elif kind == "add_chart":
                s = sheet(op["sheet"])
                s["charts"].append(op["chart"])
                changed.append(s["id"])
            elif kind == "delete_rows":
                s = sheet(op["sheet"])
                hdr = 1 if s["columns"] else 0
                start = int(op["row"]) - 1 - hdr
                del s["rows"][start:start + int(op.get("count", 1))]
                changed.append(s["id"])
            elif kind == "insert_rows":
                s = sheet(op["sheet"])
                hdr = 1 if s["columns"] else 0
                start = int(op["row"]) - 1 - hdr
                for k, rowv in enumerate(op.get("values") or [[]]):
                    s["rows"].insert(start + k, rowv)
                changed.append(s["id"])
        else:
            raise OpError(f"不支持的操作：{kind}")
    if not d["sheets"]:
        raise OpError("至少保留一个工作表")
    new = Workbook.model_validate(d)
    return new, list(dict.fromkeys(changed))


# ---------- 差异 ----------

def deck_diff(a: Deck, b: Deck) -> dict:
    """两个版本之间：新增、删除、修改的页面。"""
    am = {s.id: s.model_dump() for s in a.slides}
    bm = {s.id: s.model_dump() for s in b.slides}
    return {"added": [i for i in bm if i not in am], "removed": [i for i in am if i not in bm],
            "modified": [i for i in bm if i in am and am[i] != bm[i]],
            "theme_changed": a.theme != b.theme,
            "order_changed": list(am) != list(bm),
            "settings_changed": a.model_dump(exclude={"slides", "theme"}) != b.model_dump(exclude={"slides", "theme"})}


def doc_diff(a: Document, b: Document) -> dict:
    am = {x.id: x.model_dump() for x in a.blocks}
    bm = {x.id: x.model_dump() for x in b.blocks}
    return {"added": [i for i in bm if i not in am], "removed": [i for i in am if i not in bm],
            "modified": [i for i in bm if i in am and am[i] != bm[i]]}
