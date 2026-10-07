"""AI 生成 Excel。

- 有上传数据时：模型只输出受限的操作描述（筛选、分组、汇总、排序、透视、派生列），由 pandas 执行；
  合计行写成真实公式，并把 pandas 独立计算的结果作为验证依据。不运行模型写的任意代码。
- 没有数据时：模型直接设计工作簿（表头、公式、格式、图表），生成后重算验证，发现错误修正一轮。
"""
from __future__ import annotations

import json
import math
from pathlib import Path

from openpyxl.utils import get_column_letter

from .. import storage
from ..spec.ops import apply_wb_ops
from ..spec.workbook import Workbook, workbook_text
from ..tools import limits
from ..util import UserError
from . import prompts
from .common import ensure_work, gather_sources, render, save_version, source_block
from .context import Ctx

MAX_DATA_ROWS = 20000


class _LazyPandas:
    """pandas 约占 50 MB 内存，只有分析上传数据时才用到：第一次访问时才导入，之后替换为真正的模块。
    worker 和 Web 进程导入本模块时不再连带加载 pandas。"""

    def __getattr__(self, name):
        import pandas

        globals()["pd"] = pandas
        return getattr(pandas, name)


pd = _LazyPandas()


def handle_gen_xls(ctx: Ctx):
    p = ctx.p
    frames, other_files = load_data(ctx, p.get("file_ids") or [])
    ctx.check()
    if frames:
        wb_d, summary = plan_from_data(ctx, frames)
    else:
        sources, _ = gather_sources(ctx, other_files, 0.0, 0.15)
        wb_d = generate_workbook(ctx, sources)
        summary = ""
    title = wb_d.get("title") or p.get("topic", "")[:30] or "工作簿"
    work_id = ensure_work(ctx, "xls", title)
    wb = Workbook.model_validate(wb_d)
    ctx.progress(0.6, "生成与验证")
    rendered = render(ctx, "xls", wb.model_dump(mode="json"), {}, "render_0", title=title)
    fixable = [i for i in rendered.get("issues", []) if i.get("kind") in ("error_value", "ref_out_of_range", "bad_ref", "mismatch")]
    fix_log = []
    if fixable and not frames:
        ctx.progress(0.75, "修正公式")
        try:
            wb2 = fix_formulas(ctx, wb, fixable)
            r2 = render(ctx, "xls", wb2.model_dump(mode="json"), {}, "render_1", title=title)
            if len(r2.get("issues", [])) < len(rendered.get("issues", [])):
                wb, rendered = wb2, r2
                fix_log.append(f"修正了公式问题（剩余 {len(r2.get('issues', []))} 处）")
        except Exception as e:
            fix_log.append(f"自动修正失败：{str(e)[:100]}")
    v = save_version(ctx, work_id, "xls", wb.model_dump(mode="json"), rendered, {}, source="generate",
                     message=p.get("topic", "")[:200], title=title, search_text=workbook_text(wb),
                     extra={"summary": summary, "fix_log": fix_log})
    return {"work_id": work_id, "version_id": v["id"], "issues": len(rendered.get("issues", [])), "summary": summary}


# ---------- 数据读取 ----------

def load_data(ctx: Ctx, file_ids: list[str]) -> tuple[dict[str, pd.DataFrame], list[str]]:
    frames: dict[str, pd.DataFrame] = {}
    others = []
    for fid in file_ids:
        f = storage.get_file(fid)
        if not f:
            continue
        path = ctx.materialize(f)
        kind = (f.get("meta") or {}).get("kind") or limits.sniff(path, f["name"])
        stem = Path(f["name"]).stem[:20]
        if kind == "csv":
            df = None
            for enc in ("utf-8-sig", "gb18030"):
                try:
                    df = pd.read_csv(path, encoding=enc)
                    break
                except UnicodeDecodeError:
                    continue
            if df is not None:
                frames[stem] = df
        elif kind in ("xlsx", "xlsm", "xls", "ods"):
            if kind != "xlsx":
                from ..tools import office
                path = office.normalize_to_ooxml(path, ctx.tmp / "norm", cancel=ctx.cancelled) if kind != "xlsm" else path
            for name, df in pd.read_excel(path, sheet_name=None).items():
                if not df.empty:
                    frames[f"{stem}_{name}"[:28]] = df
        else:
            others.append(fid)
    return frames, others


def _schema(frames: dict[str, pd.DataFrame]) -> str:
    out = []
    for name, df in frames.items():
        cols = ", ".join(f"{c}（{_dtype(df[c])}）" for c in df.columns)
        out.append(f"数据表 {name}：{len(df)} 行\n列：{cols}\n前 5 行：\n{df.head(5).to_csv(index=False)}")
    return "\n\n".join(out)


def _dtype(s: pd.Series) -> str:
    if pd.api.types.is_numeric_dtype(s):
        return "数字"
    if pd.api.types.is_datetime64_any_dtype(s):
        return "日期"
    return "文字"


# ---------- 有数据：受限操作 + pandas 执行 ----------

AGG = {"sum": "sum", "mean": "mean", "count": "count", "min": "min", "max": "max", "median": "median"}


def run_ops(df: pd.DataFrame, ops: list[dict]) -> pd.DataFrame:
    for op in ops:
        k = op.get("op")
        if k == "filter":
            col, cmp, val = op["column"], op.get("cmp", "=="), op.get("value")
            s = df[col]
            if cmp == "contains":
                df = df[s.astype(str).str.contains(str(val), na=False, regex=False)]
            elif cmp == "in":
                df = df[s.isin(val if isinstance(val, list) else [val])]
            else:
                ops_map = {"==": s.__eq__, "!=": s.__ne__, ">": s.__gt__, ">=": s.__ge__, "<": s.__lt__, "<=": s.__le__}
                if cmp not in ops_map:
                    raise UserError(f"不支持的比较方式：{cmp}")
                df = df[ops_map[cmp](val)]
        elif k == "group":
            aggs = op.get("aggs") or []
            spec = {a.get("as") or f"{a['column']}_{a['func']}": pd.NamedAgg(column=a["column"], aggfunc=AGG[a["func"]]) for a in aggs}
            df = df.groupby(op["by"], dropna=False).agg(**spec).reset_index()
        elif k == "sort":
            df = df.sort_values(op["by"], ascending=bool(op.get("ascending", True)))
        elif k == "top":
            df = df.head(int(op.get("n", 10)))
        elif k == "pivot":
            df = pd.pivot_table(df, index=op["index"], columns=op["columns"], values=op["values"],
                                aggfunc=AGG[op.get("func", "sum")], fill_value=0).reset_index()
            df.columns = [str(c) for c in df.columns]
        elif k == "select":
            df = df[op["columns"]]
        elif k == "derive":
            left = df[op["left"]]
            right = df[op["right"]] if isinstance(op.get("right"), str) and op["right"] in df.columns else float(op["right"])
            o = op.get("operator", "+")
            df = df.assign(**{op["name"]: {"+": left + right, "-": left - right, "*": left * right, "/": left / right}[o]})
        else:
            raise UserError(f"不支持的数据操作：{k}")
    return df


def _cell(v):
    if v is None:
        return None
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return None
    if hasattr(v, "item"):
        try:
            v = v.item()
        except Exception:
            pass
    if isinstance(v, (pd.Timestamp,)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int, float, str, bool)):
        return v
    return str(v)


def df_sheet(name: str, df: pd.DataFrame, totals: bool = False, chart: dict | None = None, note: str = "") -> dict:
    cols = [str(c) for c in df.columns]
    rows = [[_cell(v) for v in r] for r in df.itertuples(index=False, name=None)]
    n = len(rows)
    columns = []
    for c in df.columns:
        fmt = ""
        if pd.api.types.is_float_dtype(df[c]):
            fmt = "#,##0.00"
        elif pd.api.types.is_integer_dtype(df[c]):
            fmt = "#,##0"
        columns.append({"header": str(c), "number_format": fmt})
    expected = []
    if totals and n:
        tot = ["合计"] + [None] * (len(cols) - 1)
        for j, c in enumerate(df.columns):
            if j == 0:
                continue
            if pd.api.types.is_numeric_dtype(df[c]):
                L = get_column_letter(j + 1)
                tot[j] = f"=SUM({L}2:{L}{n + 1})"
                val = float(pd.to_numeric(df[c], errors="coerce").fillna(0).sum())
                expected.append({"cell": f"{L}{n + 2}", "value": val, "tolerance": max(1e-6, abs(val) * 1e-9)})
        rows.append(tot)
    charts = []
    if chart and chart.get("x") in cols and n:
        xi = cols.index(chart["x"]) + 1
        vals = []
        for y in chart.get("y") or []:
            if y in cols:
                L = get_column_letter(cols.index(y) + 1)
                vals.append(f"{L}1:{L}{n + 1}")
        if vals:
            XL = get_column_letter(xi)
            charts.append({"type": chart.get("type", "column") if chart.get("type") in ("column", "bar", "line", "pie", "area") else "column",
                           "title": chart.get("title", ""), "categories": f"{XL}2:{XL}{n + 1}", "values": vals[:1] if chart.get("type") == "pie" else vals})
    return {"name": name[:31], "columns": columns, "rows": rows, "charts": charts, "expected": expected, "note": note}


def plan_from_data(ctx: Ctx, frames: dict[str, pd.DataFrame]) -> tuple[dict, str]:
    p = ctx.p
    user = f"用户要求：{p.get('topic', '')}\n\n<资料>\n{_schema(frames)}\n</资料>"

    ctx.progress(0.2, "制定分析方案")
    plan = ctx.llm.json("planner", prompts.XLS_PLAN, user, validate=lambda d: validate_plan(frames, d), stream=True, max_tokens=4000)
    sheets = []
    for name, df in frames.items():
        cut = df.head(MAX_DATA_ROWS)
        sheets.append(df_sheet(f"数据_{name}"[:31], cut, note="原始数据" + (f"（只保留前 {MAX_DATA_ROWS} 行）" if len(df) > MAX_DATA_ROWS else "")))
    for s in plan["sheets"]:
        res = run_ops(frames[s["source"]].copy(), s.get("ops") or [])
        sheets.append(df_sheet(s.get("name") or "结果", res.head(MAX_DATA_ROWS), bool(s.get("totals")), s.get("chart"), s.get("note", "")))
    summary = plan.get("summary", "")
    if summary:
        sheets.insert(0, {"name": "分析结论", "columns": [{"header": "结论", "width": 100}], "rows": [[summary]] +
                          [[f"{s.get('name')}：{s.get('note', '')}"] for s in plan["sheets"] if s.get("note")],
                          "freeze_header": False, "autofilter": False})
    return {"title": plan.get("title") or p.get("topic", "")[:30] or "数据分析", "sheets": sheets}, summary


def validate_plan(frames: dict[str, pd.DataFrame], d) -> dict:
    """检查模型给出的分析方案：数据表存在、每一步操作都能在真实数据上执行。
    写错列名等问题由 pandas 抛出 KeyError 等各种异常，统一转成 ValueError 反馈给模型修正，而不是让任务失败。"""
    if not isinstance(d, dict) or not isinstance(d.get("sheets"), list):
        raise ValueError("需要 sheets 数组")
    for s in d["sheets"]:
        if not isinstance(s, dict):
            raise ValueError("sheets 中的每一项都要是对象")
        src = s.get("source")
        if src not in frames:
            raise ValueError(f"数据表 {src} 不存在，可用：{', '.join(frames)}")
        try:
            run_ops(frames[src].copy(), s.get("ops") or [])
        except Exception as e:  # pandas 的 KeyError / TypeError / UserError 等
            msg = e.message if isinstance(e, UserError) else f"{type(e).__name__}: {str(e)[:200]}"
            raise ValueError(f"数据表 {src} 的操作无法执行：{msg}")
    return d


# ---------- 无数据：直接设计 ----------

def generate_workbook(ctx: Ctx, sources: str) -> dict:
    p = ctx.p
    user = f"用户要求：{p.get('topic', '')}\n" + (f"其他要求：{p['extra']}\n" if p.get("extra") else "") + "\n" + source_block(sources)
    ctx.progress(0.2, "设计工作簿")
    wb = ctx.llm.json("planner", prompts.XLS_GEN, user, model=Workbook, stream=True, max_tokens=8000)
    return wb.model_dump(mode="json")


def fix_formulas(ctx: Ctx, wb: Workbook, issues: list[dict]) -> Workbook:
    user = ("验证发现以下公式问题，请输出修正操作：\n" + "\n".join(f"- {i.get('message')}" for i in issues[:30])
            + "\n\n当前工作簿：\n" + json.dumps(wb.model_dump(mode="json"), ensure_ascii=False)[:30000])

    def validate(d):
        if not isinstance(d, dict) or not isinstance(d.get("ops"), list):
            raise ValueError("需要 ops 数组")
        return apply_wb_ops(wb, d["ops"])[0]
    return ctx.llm.json("planner", prompts.EDIT_XLS, user, validate=validate, max_tokens=6000)
