"""图标：用原生形状和自由曲线组合绘制，作为一个组合插入。

原生矢量形状在 Office 与 WPS 中都能正常显示，可直接改色、缩放，不依赖 SVG 或图标字体。
"""
from __future__ import annotations

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Pt

from .theme import hex_rgb

# 每个图标是若干图元，坐标为 0..1 的相对位置
ICONS: dict[str, list[tuple]] = {
    "star": [("preset", MSO_SHAPE.STAR_5_POINT, 0.12, 0.12, 0.76, 0.76)],
    "heart": [("preset", MSO_SHAPE.HEART, 0.15, 0.18, 0.7, 0.66)],
    "lightning": [("preset", MSO_SHAPE.LIGHTNING_BOLT, 0.2, 0.12, 0.6, 0.76)],
    "gear": [("preset", MSO_SHAPE.GEAR_6, 0.14, 0.14, 0.72, 0.72)],
    "cloud": [("preset", MSO_SHAPE.CLOUD, 0.1, 0.22, 0.8, 0.56)],
    "sun": [("preset", MSO_SHAPE.SUN, 0.1, 0.1, 0.8, 0.8)],
    "cube": [("preset", MSO_SHAPE.CUBE, 0.18, 0.18, 0.64, 0.64)],
    "diamond": [("preset", MSO_SHAPE.DIAMOND, 0.15, 0.12, 0.7, 0.76)],
    "database": [("preset", MSO_SHAPE.CAN, 0.24, 0.12, 0.52, 0.76)],
    "plus": [("preset", MSO_SHAPE.MATH_PLUS, 0.12, 0.12, 0.76, 0.76)],
    "cross": [("line", 0.25, 0.25, 0.75, 0.75, 0.1), ("line", 0.75, 0.25, 0.25, 0.75, 0.1)],
    "check": [("path", [(0.2, 0.52), (0.42, 0.74), (0.82, 0.28)], False, 0.1)],
    "document": [("rect", 0.24, 0.12, 0.52, 0.76, 0.06, "line"), ("line", 0.34, 0.36, 0.66, 0.36, 0.05),
                 ("line", 0.34, 0.52, 0.66, 0.52, 0.05), ("line", 0.34, 0.68, 0.56, 0.68, 0.05)],
    "target": [("circle", 0.5, 0.5, 0.4, "line"), ("circle", 0.5, 0.5, 0.25, "line"), ("circle", 0.5, 0.5, 0.1, "fill")],
    "growth": [("rect", 0.15, 0.58, 0.16, 0.28, 0.02, "fill"), ("rect", 0.42, 0.4, 0.16, 0.46, 0.02, "fill"),
               ("rect", 0.69, 0.16, 0.16, 0.7, 0.02, "fill")],
    "decline": [("rect", 0.15, 0.16, 0.16, 0.7, 0.02, "fill"), ("rect", 0.42, 0.4, 0.16, 0.46, 0.02, "fill"),
                ("rect", 0.69, 0.58, 0.16, 0.28, 0.02, "fill")],
    "chart": [("line", 0.14, 0.86, 0.88, 0.86, 0.05), ("line", 0.14, 0.14, 0.14, 0.86, 0.05),
              ("path", [(0.22, 0.7), (0.42, 0.48), (0.58, 0.6), (0.84, 0.26)], False, 0.07)],
    "clock": [("circle", 0.5, 0.5, 0.38, "line"), ("line", 0.5, 0.5, 0.5, 0.26, 0.07), ("line", 0.5, 0.5, 0.68, 0.6, 0.07)],
    "people": [("circle", 0.36, 0.32, 0.13, "fill"), ("rect", 0.18, 0.5, 0.36, 0.36, 0.18, "fill"),
               ("circle", 0.66, 0.36, 0.11, "fill"), ("rect", 0.54, 0.54, 0.3, 0.32, 0.15, "fill")],
    "lock": [("circle", 0.5, 0.36, 0.18, "line"), ("rect", 0.24, 0.42, 0.52, 0.44, 0.06, "fill")],
    "globe": [("circle", 0.5, 0.5, 0.38, "line"), ("ellipse", 0.5, 0.5, 0.16, 0.38, "line"), ("line", 0.12, 0.5, 0.88, 0.5, 0.05)],
    "mail": [("rect", 0.12, 0.24, 0.76, 0.52, 0.04, "line"), ("path", [(0.14, 0.27), (0.5, 0.54), (0.86, 0.27)], False, 0.05)],
    "search": [("circle", 0.42, 0.42, 0.24, "line"), ("line", 0.6, 0.6, 0.84, 0.84, 0.1)],
    "home": [("path", [(0.5, 0.14), (0.86, 0.46), (0.76, 0.46), (0.76, 0.86), (0.24, 0.86), (0.24, 0.46), (0.14, 0.46)], True, 0)],
    "flag": [("line", 0.24, 0.12, 0.24, 0.88, 0.07), ("path", [(0.27, 0.16), (0.8, 0.3), (0.27, 0.46)], True, 0)],
    "idea": [("circle", 0.5, 0.4, 0.26, "fill"), ("rect", 0.39, 0.66, 0.22, 0.2, 0.04, "fill")],
    "shield": [("path", [(0.5, 0.1), (0.84, 0.22), (0.8, 0.56), (0.5, 0.9), (0.2, 0.56), (0.16, 0.22)], True, 0)],
    "money": [("circle", 0.5, 0.5, 0.38, "line"), ("text", "¥", 0.5, 0.5)],
    "moon": [("preset", MSO_SHAPE.MOON, 0.26, 0.12, 0.46, 0.76)],
}
ICON_NAMES = sorted(ICONS)


def add_icon(shapes, name: str, x: int, y: int, size: int, color: str, badge: str | None = None, number: str | None = None):
    """在 (x, y) 处绘制 size×size 的图标。badge 为底色圆的颜色；无效名称时画编号或圆点。"""
    grp = shapes.add_group_shape()
    rgb = RGBColor(*hex_rgb(color))
    if badge:
        b = grp.shapes.add_shape(MSO_SHAPE.OVAL, x, y, size, size)
        _fill(b, badge)
        b.line.fill.background()
        pad = int(size * 0.22)
        x, y, size = x + pad, y + pad, size - 2 * pad
    prims = ICONS.get((name or "").lower())
    if prims is None:
        tb = grp.shapes.add_textbox(x, y, size, size)
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = number or "●"
        r.font.size = Pt(max(8, int(size / 12700 * 0.55)))
        r.font.bold = True
        r.font.color.rgb = rgb
        return grp
    for pr in prims:
        kind = pr[0]
        if kind == "preset":
            _, shp, px, py, pw, ph = pr
            s = grp.shapes.add_shape(shp, x + int(px * size), y + int(py * size), int(pw * size), int(ph * size))
            _fill(s, color)
            s.line.fill.background()
        elif kind in ("circle", "ellipse"):
            if kind == "circle":
                _, cx, cy, r, mode = pr
                rx = ry = r
            else:
                _, cx, cy, rx, ry, mode = pr
            s = grp.shapes.add_shape(MSO_SHAPE.OVAL, x + int((cx - rx) * size), y + int((cy - ry) * size), int(2 * rx * size), int(2 * ry * size))
            _mode(s, mode, color, size)
        elif kind == "rect":
            _, px, py, pw, ph, rad, mode = pr
            shp = MSO_SHAPE.ROUNDED_RECTANGLE if rad else MSO_SHAPE.RECTANGLE
            s = grp.shapes.add_shape(shp, x + int(px * size), y + int(py * size), int(pw * size), int(ph * size))
            if rad:
                s.adjustments[0] = min(0.5, rad / max(0.01, min(pw, ph)))
            _mode(s, mode, color, size)
        elif kind == "line":
            _, x1, y1, x2, y2, wdt = pr
            ln = grp.shapes.add_connector(1, x + int(x1 * size), y + int(y1 * size), x + int(x2 * size), y + int(y2 * size))
            _nostyle(ln)
            ln.line.color.rgb = rgb
            ln.line.width = Emu(max(6350, int(wdt * size)))
            _round_cap(ln)
        elif kind == "path":
            _, pts, closed, wdt = pr
            fb = grp.shapes.build_freeform(x + int(pts[0][0] * size), y + int(pts[0][1] * size), scale=1.0)
            fb.add_line_segments([(x + int(px * size), y + int(py * size)) for px, py in pts[1:]], close=closed)
            s = fb.convert_to_shape()
            if closed:
                _fill(s, color)
                s.line.fill.background()
            else:
                _nostyle(s)
                s.fill.background()
                s.line.color.rgb = rgb
                s.line.width = Emu(max(6350, int(wdt * size)))
                _round_cap(s)
        elif kind == "text":
            _, txt, cx, cy = pr
            tb = grp.shapes.add_textbox(x, y, size, size)
            tf = tb.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(max(8, int(size / 12700 * 0.5)))
            r.font.bold = True
            r.font.color.rgb = rgb
    return grp


def _nostyle(shape) -> None:
    from pptx.oxml.ns import qn
    el = shape._element.find(qn("p:style"))
    if el is not None:
        shape._element.remove(el)


def _fill(s, color: str) -> None:
    _nostyle(s)
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor(*hex_rgb(color))


def _mode(s, mode: str, color: str, size: int) -> None:
    _nostyle(s)
    if mode == "fill":
        _fill(s, color)
        s.line.fill.background()
    else:
        s.fill.background()
        s.line.color.rgb = RGBColor(*hex_rgb(color))
        s.line.width = Emu(max(6350, int(0.06 * size)))


def _round_cap(shape) -> None:
    from pptx.oxml.ns import qn

    ln = shape.line._get_or_add_ln()
    ln.set("cap", "rnd")
    _ = qn
