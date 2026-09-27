"""XLSX 渲染器：规格 → Excel 工作簿（真实公式、数字格式、条件格式、数据验证、原生图表）。

文件设置为打开时重新计算。验证由 app.tools.xlsx_verify 在副本上用 LibreOffice 重算完成。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook as XWorkbook
from openpyxl.chart import AreaChart, BarChart, LineChart, PieChart, Reference, ScatterChart, Series
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

from ..spec.common import Theme
from ..spec.workbook import Workbook


@dataclass
class XlsxRenderResult:
    path: Path
    sheets: list[dict] = field(default_factory=list)   # [{id, name, rows, cols}]
    counts: dict = field(default_factory=dict)         # 图表、条件格式、数据验证数量，供验证比对


def render_workbook(wb: Workbook, out: Path, theme: Theme | None = None) -> XlsxRenderResult:
    theme = theme or Theme()
    x = XWorkbook()
    x.remove(x.active)
    head_fill = PatternFill("solid", fgColor=theme.primary.lstrip("#"))
    head_font = Font(bold=True, color="FFFFFF", name="微软雅黑")
    body_font = Font(name="微软雅黑", size=11)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    res = XlsxRenderResult(path=out)
    counts = {"charts": 0, "cond_formats": 0, "validations": 0}
    for sh in wb.sheets:
        ws = x.create_sheet(sh.name)
        r0 = 1
        if sh.columns:
            for c, col in enumerate(sh.columns, 1):
                cell = ws.cell(row=1, column=c, value=col.header)
                cell.fill = head_fill
                cell.font = head_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = border
            r0 = 2
        for ri, row in enumerate(sh.rows):
            for ci, v in enumerate(row, 1):
                cell = ws.cell(row=r0 + ri, column=ci, value=v)
                cell.font = body_font
                cell.border = border
                if ci <= len(sh.columns) and sh.columns[ci - 1].number_format:
                    cell.number_format = sh.columns[ci - 1].number_format
                if isinstance(v, str) and not v.startswith("="):
                    cell.alignment = Alignment(vertical="center", wrap_text=len(v) > 30)
        nrows, ncols = sh.data_range
        # 汇总行（首列含“合计/总计”）加粗
        for ri, row in enumerate(sh.rows):
            if row and isinstance(row[0], str) and row[0].strip() in ("合计", "总计", "小计", "Total", "平均"):
                for ci in range(1, ncols + 1):
                    ws.cell(row=r0 + ri, column=ci).font = Font(name="微软雅黑", size=11, bold=True)
        # 列宽
        for ci in range(1, ncols + 1):
            if ci <= len(sh.columns) and sh.columns[ci - 1].width:
                w = sh.columns[ci - 1].width
            else:
                vals = [sh.columns[ci - 1].header] if ci <= len(sh.columns) else []
                vals += [r[ci - 1] for r in sh.rows[:200] if ci - 1 < len(r) and r[ci - 1] is not None and not (isinstance(r[ci - 1], str) and r[ci - 1].startswith("="))]
                w = min(50, max(8, max((_disp(str(v)) for v in vals), default=8) + 2))
            ws.column_dimensions[get_column_letter(ci)].width = w
        if sh.columns and sh.freeze_header:
            ws.freeze_panes = "A2"
        if sh.columns and sh.autofilter and nrows > 1:
            ws.auto_filter.ref = f"A1:{get_column_letter(ncols)}{nrows}"
        for cf in sh.cond_formats:
            color = cf.color.lstrip("#")
            if cf.kind == "color_scale":
                rule = ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B")
            elif cf.kind == "data_bar":
                rule = DataBarRule(start_type="min", end_type="max", color=theme.secondary.lstrip("#"))
            else:
                op = {"greater_than": "greaterThan", "less_than": "lessThan", "between": "between", "equal": "equal"}[cf.kind]
                formula = [str(cf.value if cf.value is not None else 0)]
                if cf.kind == "between":
                    formula.append(str(cf.value2 if cf.value2 is not None else 0))
                rule = CellIsRule(operator=op, formula=formula, fill=PatternFill("solid", fgColor=color))
            ws.conditional_formatting.add(cf.range, rule)
            counts["cond_formats"] += 1
        for v in sh.validations:
            dv = DataValidation(type="list", formula1='"' + ",".join(o.replace(",", "，") for o in v.options) + '"', allow_blank=True)
            dv.add(v.range)
            ws.add_data_validation(dv)
            counts["validations"] += 1
        for ch in sh.charts:
            chart = _chart(ws, ch, theme)
            anchor = ch.anchor or f"{get_column_letter(ncols + 2)}2"
            ws.add_chart(chart, anchor)
            counts["charts"] += 1
        res.sheets.append({"id": sh.id, "name": sh.name, "rows": nrows, "cols": ncols})
    x.calculation = CalcProperties(fullCalcOnLoad=True)
    x.properties.title = wb.title
    x.properties.creator = "DocWork"
    x.save(out)
    res.counts = counts
    return res


def _chart(ws, ch, theme: Theme):
    kinds = {"column": BarChart, "bar": BarChart, "line": LineChart, "pie": PieChart, "area": AreaChart, "scatter": ScatterChart}
    c = kinds[ch.type]()
    if ch.type == "column":
        c.type = "col"
    elif ch.type == "bar":
        c.type = "bar"
    c.title = ch.title or None
    c.height, c.width = 8, 16
    c1, r1, c2, r2 = range_boundaries(ch.categories)
    cats = Reference(ws, min_col=c1, min_row=r1, max_col=c2, max_row=r2)
    if ch.type == "scatter":
        for rng in ch.values:
            a1, b1, a2, b2 = range_boundaries(rng)
            vals = Reference(ws, min_col=a1, min_row=b1 + 1, max_col=a2, max_row=b2)
            s = Series(vals, cats, title=str(ws.cell(row=b1, column=a1).value or ""))
            s.marker.symbol = "circle"
            s.graphicalProperties.line.noFill = True
            c.series.append(s)
    else:
        for rng in ch.values:
            a1, b1, a2, b2 = range_boundaries(rng)
            data = Reference(ws, min_col=a1, min_row=b1, max_col=a2, max_row=b2)
            c.add_data(data, titles_from_data=True)
        c.set_categories(cats)
    if ch.y_title and hasattr(c, "y_axis"):
        c.y_axis.title = ch.y_title
    colors = [theme.primary, theme.secondary, theme.accent, theme.muted]
    if ch.type != "pie":
        for i, s in enumerate(c.series):
            col = colors[i % len(colors)].lstrip("#")
            s.graphicalProperties.solidFill = col if ch.type in ("column", "bar", "area") else None
            s.graphicalProperties.line.solidFill = col
    # openpyxl 3.1 默认隐藏坐标轴，显式打开
    for ax in ("x_axis", "y_axis"):
        if hasattr(c, ax):
            try:
                getattr(c, ax).delete = False
            except Exception:
                pass
    return c


def _disp(s: str) -> int:
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)
