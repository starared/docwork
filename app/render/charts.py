"""图表：PPT 原生图表（可编辑数据）、matplotlib 图片图表、图表数据 XLSX。"""
from __future__ import annotations

import io
from functools import lru_cache
from pathlib import Path

from ..spec.common import ChartSpec, Theme
from .theme import hex_rgb, series_colors


# ---------- 原生 PPT 图表 ----------

def add_native_chart(slide, spec: ChartSpec, x, y, w, h, theme: Theme, font_size: float = 12):
    from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
    from pptx.util import Pt

    t = spec.type
    stacked = spec.stacked
    types = {
        "column": XL_CHART_TYPE.COLUMN_STACKED if stacked else XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_STACKED if stacked else XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE_MARKERS,
        "area": XL_CHART_TYPE.AREA_STACKED if stacked else XL_CHART_TYPE.AREA,
        "pie": XL_CHART_TYPE.PIE,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
        "scatter": XL_CHART_TYPE.XY_SCATTER,
        "radar": XL_CHART_TYPE.RADAR_MARKERS,
        "bubble": XL_CHART_TYPE.BUBBLE,
    }
    fmt = spec.number_format or "General"
    if t in ("scatter", "bubble"):
        data = BubbleChartData() if t == "bubble" else XyChartData()
        for ps in spec.point_series:
            s = data.add_series(ps.name or "系列")
            for p in ps.points:
                if t == "bubble":
                    s.add_data_point(p.x, p.y, p.size if p.size is not None else 1)
                else:
                    s.add_data_point(p.x, p.y)
    else:
        data = CategoryChartData(number_format=fmt)
        # 条形图自下而上绘制：数据倒序写入，使显示顺序与规格一致（不改坐标轴交叉设置，兼容性更好）
        rev = t == "bar"
        data.categories = list(reversed(spec.categories)) if rev else spec.categories
        for s in (reversed(spec.series) if rev else spec.series):
            vals = [v if v is not None else None for v in s.values]
            data.add_series(s.name or "系列", list(reversed(vals)) if rev else vals)
    gf = slide.shapes.add_chart(types[t], x, y, w, h, data)
    chart = gf.chart
    chart.font.size = Pt(font_size)
    chart.font.name = theme.body_font
    _set_ea_font(chart, theme.body_font)
    chart.font.color.rgb = RGBColor(*hex_rgb(theme.text))
    nser = len(spec.point_series) if t in ("scatter", "bubble") else len(spec.series)
    single_pie = t in ("pie", "doughnut")
    chart.has_legend = single_pie or nser > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(font_size)
    if spec.title:
        chart.has_title = True
        chart.chart_title.text_frame.text = spec.title
        p = chart.chart_title.text_frame.paragraphs[0]
        p.runs[0].font.size = Pt(font_size + 2)
        p.runs[0].font.bold = True
    else:
        chart.has_title = False
    colors = series_colors(theme, max(nser, len(spec.categories) if single_pie else nser))
    plot = chart.plots[0]
    if single_pie:
        pts = plot.series[0].points
        for i in range(len(spec.categories)):
            f = pts[i].format.fill
            f.solid()
            f.fore_color.rgb = RGBColor(*hex_rgb(colors[i % len(colors)]))
    else:
        for i, s in enumerate(plot.series):
            c = RGBColor(*hex_rgb(colors[(nser - 1 - i) if t == "bar" else i]))
            if t in ("line", "radar", "scatter"):
                s.format.line.color.rgb = c
                s.format.line.width = Pt(2.25)
                try:
                    s.marker.format.fill.solid()
                    s.marker.format.fill.fore_color.rgb = c
                    s.marker.format.line.color.rgb = c
                except Exception:
                    pass
                if t == "scatter":
                    s.format.line.fill.background()
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = c
        if t in ("column", "bar"):
            try:
                plot.gap_width = 80
                if stacked:
                    plot.overlap = 100
            except Exception:
                pass
    npts = len(spec.categories) * max(1, nser)
    if spec.show_values and (single_pie or npts <= 12) and t not in ("scatter", "bubble", "radar", "area"):
        plot.has_data_labels = True
        dl = plot.data_labels
        dl.font.size = Pt(max(9, font_size - 1))
        if single_pie:
            dl.number_format = spec.number_format or "0%"
            dl.number_format_is_linked = False
            dl.show_percentage = not spec.number_format
            dl.show_value = bool(spec.number_format)
            dl.position = XL_LABEL_POSITION.OUTSIDE_END if t == "pie" else XL_LABEL_POSITION.CENTER
            if t == "doughnut":
                # 环形图标签在色块内：按色块颜色选择黑或白字
                from .theme import on_color
                for i in range(len(spec.categories)):
                    pdl = plot.series[0].points[i].data_label
                    pdl.font.size = Pt(max(9, font_size - 1))
                    pdl.font.bold = True
                    pdl.font.color.rgb = RGBColor(*hex_rgb(on_color(colors[i % len(colors)])))
                    _dlbl_flags(pdl, show_val=bool(spec.number_format), show_pct=not spec.number_format,
                                fmt=spec.number_format or "0%")
        else:
            if spec.number_format:
                dl.number_format = spec.number_format
                dl.number_format_is_linked = False
            dl.show_value = True
            if not stacked and t in ("column", "bar"):
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
    if not single_pie and t != "radar":
        try:
            va = chart.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = RGBColor(0xE5, 0xE7, 0xEB)
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(font_size - 1)
            if spec.number_format:
                va.tick_labels.number_format = spec.number_format
                va.tick_labels.number_format_is_linked = False
            ca = chart.category_axis
            ca.tick_labels.font.size = Pt(font_size - 1)
            ca.format.line.color.rgb = RGBColor(0xD1, 0xD5, 0xDB)
            ca.has_major_gridlines = False
        except Exception:
            pass
    return gf


def native_chart_parts(spec: ChartSpec, theme: Theme, width_emu: int, height_emu: int, font_size: float = 10) -> tuple[bytes, bytes, str]:
    """生成与 PPT 完全相同样式的原生图表，供 Word 使用：返回（图表 XML、内嵌数据工作簿、XML 中引用工作簿的关系 ID）。
    做法是在一个临时演示文稿里用 add_native_chart 画图，再取出图表部件；Word 与 PowerPoint 的图表部件格式相同。"""
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as PRT

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    gf = add_native_chart(slide, spec, 0, 0, width_emu, height_emu, theme, font_size)
    cpart = gf.chart.part
    xlsx_rel = next(r for r in cpart.rels.values() if r.reltype == PRT.PACKAGE)
    return cpart.blob, xlsx_rel.target_part.blob, xlsx_rel.rId


def _dlbl_flags(pdl, show_val: bool, show_pct: bool, fmt: str) -> None:
    """单个数据点标签的显示项和数字格式（python-pptx 未提供，直接写 XML，注意元素顺序）。"""
    from pptx.oxml.ns import qn

    d = pdl._get_or_add_dLbl()
    for tag, val in (("c:showVal", show_val), ("c:showPercent", show_pct)):
        el = d.find(qn(tag))
        if el is not None:
            el.set("val", "1" if val else "0")
    nf = d.find(qn("c:numFmt"))
    if nf is None:
        nf = d.makeelement(qn("c:numFmt"), {})
        idx = d.find(qn("c:idx"))
        idx.addnext(nf)
    nf.set("formatCode", fmt)
    nf.set("sourceLinked", "0")


def _set_ea_font(chart, family: str) -> None:
    """图表文字的东亚字体。"""
    from pptx.oxml.ns import qn

    txPr = chart._chartSpace.chart.getparent().find(qn("c:txPr"))
    if txPr is None:
        return
    for rpr in txPr.iter(qn("a:defRPr")):
        ea = rpr.find(qn("a:ea"))
        if ea is None:
            ea = rpr.makeelement(qn("a:ea"), {})
            rpr.append(ea)
        ea.set("typeface", family)


# ---------- matplotlib 图片图表 ----------

@lru_cache(maxsize=1)
def _mpl():
    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager

    from .fonts import font_file

    for fam in ("微软雅黑", "宋体"):
        f = font_file(fam)
        if f:
            try:
                font_manager.fontManager.addfont(f)
            except Exception:
                pass
    names = [fe.name for fe in font_manager.fontManager.ttflist if "CJK" in fe.name]
    plt.rcParams["font.sans-serif"] = (names[:1] or []) + ["DejaVu Sans"]
    plt.rcParams["font.family"] = "sans-serif"
    plt.rcParams["axes.unicode_minus"] = False
    return plt


import threading

MPL_LOCK = threading.RLock()  # pyplot 不是线程安全的


def chart_png(spec: ChartSpec, theme: Theme, width_in: float = 8, height_in: float = 4.5, dpi: int = 200) -> bytes:
    with MPL_LOCK:
        return _chart_png(spec, theme, width_in, height_in, dpi)


def _chart_png(spec: ChartSpec, theme: Theme, width_in: float, height_in: float, dpi: int) -> bytes:
    plt = _mpl()
    import numpy as np

    plt.rcParams["font.size"] = 13
    fig, ax = plt.subplots(figsize=(width_in, height_in), dpi=dpi)
    fig.patch.set_facecolor(theme.background)
    ax.set_facecolor(theme.background)
    tc = theme.text
    colors = series_colors(theme, max(len(spec.series), len(spec.categories), 3))
    t = spec.type
    cats = spec.categories
    if t in ("column", "bar", "line", "area"):
        x = np.arange(len(cats))
        n = len(spec.series)
        bw = 0.8 / max(1, n) if not spec.stacked else 0.6
        bottom = np.zeros(len(cats))
        for i, s in enumerate(spec.series):
            vals = np.array([v or 0 for v in s.values], dtype=float)
            if t == "column":
                if spec.stacked:
                    ax.bar(x, vals, bw, bottom=bottom, label=s.name, color=colors[i])
                    bottom += vals
                else:
                    ax.bar(x - 0.4 + bw * (i + 0.5), vals, bw, label=s.name, color=colors[i])
            elif t == "bar":
                if spec.stacked:
                    ax.barh(x, vals, bw, left=bottom, label=s.name, color=colors[i])
                    bottom += vals
                else:
                    ax.barh(x - 0.4 + bw * (i + 0.5), vals, bw, label=s.name, color=colors[i])
            elif t == "line":
                ax.plot(x, vals, marker="o", label=s.name, color=colors[i], linewidth=2)
            else:
                if spec.stacked:
                    ax.fill_between(x, bottom, bottom + vals, label=s.name, color=colors[i], alpha=0.85)
                    bottom += vals
                else:
                    ax.fill_between(x, vals, label=s.name, color=colors[i], alpha=0.5)
        if t == "bar":
            ax.set_yticks(x, cats)
            ax.invert_yaxis()
        else:
            ax.set_xticks(x, cats, rotation=0 if len(cats) <= 8 else 30)
    elif t in ("pie", "doughnut"):
        vals = [v or 0 for v in spec.series[0].values]
        wedges, *_ = ax.pie(vals, labels=cats, colors=colors[: len(vals)], autopct="%1.0f%%", startangle=90,
                            wedgeprops={"width": 0.45} if t == "doughnut" else None, textprops={"color": tc})
        ax.axis("equal")
    elif t in ("scatter", "bubble"):
        for i, ps in enumerate(spec.point_series):
            xs = [p.x for p in ps.points]
            ys = [p.y for p in ps.points]
            ss = [max(10, (p.size or 1) * 30) for p in ps.points] if t == "bubble" else 40
            ax.scatter(xs, ys, s=ss, label=ps.name, color=colors[i], alpha=0.75)
    elif t == "radar":
        ang = np.linspace(0, 2 * np.pi, len(cats), endpoint=False).tolist()
        fig.clf()
        ax = fig.add_subplot(111, polar=True)
        for i, s in enumerate(spec.series):
            vals = [v or 0 for v in s.values]
            ax.plot(ang + ang[:1], vals + vals[:1], color=colors[i], label=s.name)
            ax.fill(ang + ang[:1], vals + vals[:1], color=colors[i], alpha=0.15)
        ax.set_xticks(ang, cats)
    elif t == "waterfall":
        vals = [v or 0 for v in spec.series[0].values]
        cum = 0.0
        for i, v in enumerate(vals):
            color = theme.secondary if v >= 0 else theme.accent
            ax.bar(i, v, bottom=cum if v >= 0 else cum + v, color=color, width=0.6)
            ax.text(i, cum + max(v, 0), f"{v:+g}", ha="center", va="bottom", color=tc, fontsize=12)
            cum += v
        ax.bar(len(vals), cum, color=theme.primary, width=0.6)
        ax.text(len(vals), cum, f"{cum:g}", ha="center", va="bottom", color=tc, fontsize=12)
        ax.set_xticks(range(len(vals) + 1), list(cats) + ["合计"])
    elif t == "funnel":
        vals = [v or 0 for v in spec.series[0].values]
        mx = max(vals) or 1
        for i, v in enumerate(vals):
            ax.barh(i, v, left=(mx - v) / 2, color=colors[i % len(colors)], height=0.8)
            ax.text(mx / 2, i, f"{cats[i]}  {v:g}", ha="center", va="center", color="white", fontsize=10)
        ax.invert_yaxis()
        ax.axis("off")
    elif t == "histogram":
        vals = [v for s in spec.series for v in s.values if v is not None]
        ax.hist(vals, bins=min(20, max(5, len(vals) // 3)), color=theme.primary)
    elif t == "boxplot":
        ax.boxplot([[v for v in s.values if v is not None] for s in spec.series], tick_labels=[s.name for s in spec.series])
    elif t == "heatmap":
        arr = np.array([[v or 0 for v in s.values] for s in spec.series])
        im = ax.imshow(arr, cmap="Blues")
        ax.set_xticks(range(len(cats)), cats)
        ax.set_yticks(range(len(spec.series)), [s.name for s in spec.series])
        fig.colorbar(im, ax=ax)
    elif t == "sankey":
        _sankey(ax, spec, colors, tc)
    if spec.title:
        ax.set_title(spec.title, color=tc, fontsize=13, fontweight="bold")
    if t not in ("pie", "doughnut", "funnel", "sankey", "heatmap"):
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        ax.tick_params(colors=tc)
        for side in ("left", "bottom"):
            ax.spines[side].set_color("#9CA3AF")
        if spec.unit:
            ax.set_ylabel(spec.unit, color=tc)
        ax.grid(axis="x" if t == "bar" else "y", color="#E5E7EB", linewidth=0.8)
        ax.set_axisbelow(True)
        if (len(spec.series) > 1 or len(spec.point_series) > 1) and ax.get_legend_handles_labels()[0]:
            ax.legend(frameon=False, labelcolor=tc)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor=fig.get_facecolor())
    plt.close(fig)
    return buf.getvalue()


def _sankey(ax, spec: ChartSpec, colors, tc):
    """简化桑基图：两列或多列节点，按流量宽度绘制带状连线。"""
    import numpy as np
    from matplotlib.patches import PathPatch
    from matplotlib.path import Path as MPath

    nodes: list[str] = []
    for f in spec.flows:
        for n in (f.source, f.target):
            if n not in nodes:
                nodes.append(n)
    # 层级：没有入边的在第 0 层
    level = {n: 0 for n in nodes}
    for _ in range(len(nodes)):
        for f in spec.flows:
            level[f.target] = max(level[f.target], level[f.source] + 1)
    cols: dict[int, list[str]] = {}
    for n in nodes:
        cols.setdefault(level[n], []).append(n)
    size = {n: max(sum(f.value for f in spec.flows if f.source == n), sum(f.value for f in spec.flows if f.target == n)) for n in nodes}
    total = max(sum(size[n] for n in c) for c in cols.values()) or 1
    gap = 0.03
    pos = {}
    for lv, ns in cols.items():
        y = 0.0
        for n in ns:
            hgt = size[n] / total * (1 - gap * (len(ns) - 1))
            pos[n] = [lv, y, hgt, y, y]  # level, top, height, out cursor, in cursor
            y += hgt + gap
    maxlv = max(cols) or 1
    w = 0.02
    for i, n in enumerate(nodes):
        lv, y0, hh, _, _ = pos[n]
        x0 = lv / maxlv
        ax.add_patch(__import__("matplotlib.patches", fromlist=["Rectangle"]).Rectangle((x0 - w / 2, y0), w, hh, color=colors[i % len(colors)]))
        ax.text(x0 + (w if lv < maxlv else -w), y0 + hh / 2, n, va="center", ha="left" if lv < maxlv else "right", color=tc, fontsize=9)
    for f in spec.flows:
        s, t = pos[f.source], pos[f.target]
        hh = f.value / total * (1 - gap * 2)
        xs, xt = s[0] / maxlv + w / 2, t[0] / maxlv - w / 2
        ys, yt = s[3], t[4]
        s[3] += hh
        t[4] += hh
        mid = (xs + xt) / 2
        verts = [(xs, ys), (mid, ys), (mid, yt), (xt, yt), (xt, yt + hh), (mid, yt + hh), (mid, ys + hh), (xs, ys + hh), (xs, ys)]
        codes = [MPath.MOVETO, MPath.CURVE4, MPath.CURVE4, MPath.CURVE4, MPath.LINETO, MPath.CURVE4, MPath.CURVE4, MPath.CURVE4, MPath.CLOSEPOLY]
        ax.add_patch(PathPatch(MPath(verts, codes), alpha=0.35, color=colors[nodes.index(f.source) % len(colors)], lw=0))
    ax.set_xlim(-0.05, 1.05)
    ax.set_ylim(1.02, -0.02)
    ax.axis("off")


# ---------- 图表数据 XLSX ----------

def chart_data_xlsx(charts: list[tuple[str, ChartSpec]], path: Path) -> Path:
    """把图表数据写成 XLSX 附件：每个图表一个工作表。"""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for i, (label, c) in enumerate(charts):
        ws = wb.create_sheet(f"{i + 1}_{label}"[:31].replace("/", "_").replace(":", "_"))
        ws.append([c.title or label])
        if c.type in ("scatter", "bubble"):
            ws.append(["系列", "x", "y", "size"])
            for ps in c.point_series:
                for p in ps.points:
                    ws.append([ps.name, p.x, p.y, p.size])
        elif c.type == "sankey":
            ws.append(["来源", "去向", "数值"])
            for f in c.flows:
                ws.append([f.source, f.target, f.value])
        else:
            ws.append(["分类"] + [s.name for s in c.series])
            for j, cat in enumerate(c.categories or [str(k + 1) for k in range(len(c.series[0].values))]):
                ws.append([cat] + [s.values[j] if j < len(s.values) else None for s in c.series])
        if c.source:
            ws.append([])
            ws.append(["数据来源", c.source])
    if not wb.sheetnames:
        wb.create_sheet("数据")
    wb.save(path)
    return path
