"""PPTX 渲染器：规格 → 原生可编辑的 PowerPoint 文件。

布局用代码按网格计算坐标；文字按字号自动适配（正文不低于 14pt），并记录每个元素的
位置、预期文字和是否预估溢出，供排版检查和网页上的点击选中使用。
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from ..spec.common import ChartSpec, ImageRef, TableSpec
from ..spec.deck import Deck, Slide
from .charts import add_native_chart, chart_png
from .fonts import fit_text, measurer
from .icons import add_icon
from .theme import contrast, ensure_contrast, heading_color, hex_rgb, mix, normalize_theme, on_color

EMU_IN = 914400
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
INSET_X, INSET_Y = 0.05, 0.03  # 英寸
BODY_MIN = 14
SMALL_MIN = 11


@dataclass
class Elem:
    id: str
    slide: str
    kind: str
    box: list[float]           # 以页面宽高为 1 的相对坐标 [x, y, w, h]
    text: str = ""
    size: float = 0
    min_size: float = 0
    overflow: bool = False
    role: str = ""
    decor: bool = False


@dataclass
class RenderResult:
    path: Path
    elements: dict[str, list[dict]] = field(default_factory=dict)
    issues: list[dict] = field(default_factory=list)
    image_charts: list[tuple[str, ChartSpec]] = field(default_factory=list)
    all_charts: list[tuple[str, ChartSpec]] = field(default_factory=list)
    slide_ids: list[str] = field(default_factory=list)


def _rgb(h: str) -> RGBColor:
    return RGBColor(*hex_rgb(h))


def E(v: float) -> Emu:
    return Emu(int(round(v * EMU_IN)))


_MD = re.compile(r"\*\*(.+?)\*\*")


class SlideRenderer:
    def __init__(self, deck: Deck, slide, s: Slide, W: float, H: float, index: int, total: int,
                 assets: dict[str, Path], tmp: Path, result: RenderResult):
        self.deck, self.slide, self.s = deck, slide, s
        self.W, self.H = W, H
        self.t = deck.theme
        self.index, self.total = index, total
        self.assets, self.tmp, self.result = assets, tmp, result
        self.elems: list[Elem] = []
        self.mx = 0.6 if W > 11 else 0.5
        self.cy = 1.5
        self.cb = H - 0.6
        self.cw = W - 2 * self.mx
        self.ch = self.cb - self.cy
        self.bg = self.t.background

    # ----- 基础图元 -----

    def rel(self, x, y, w, h) -> list[float]:
        return [round(x / self.W, 5), round(y / self.H, 5), round(w / self.W, 5), round(h / self.H, 5)]

    def rect(self, x, y, w, h, fill: str | None, line: str | None = None, shape=MSO_SHAPE.RECTANGLE, radius: float | None = None,
             alpha: float | None = None, line_w: float = 1.0, name: str = ""):
        sh = self.slide.shapes.add_shape(shape, E(x), E(y), E(w), E(h))
        if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
            sh.adjustments[0] = max(0.0, min(0.5, radius / max(0.01, min(w, h))))
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = _rgb(fill)
            if alpha is not None:
                _set_alpha(sh, alpha)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = _rgb(line)
            sh.line.width = Pt(line_w)
        else:
            sh.line.fill.background()
        _nostyle(sh)
        if name:
            sh.name = name
        # 形状默认带文字格式，清空以免影响
        return sh

    def line(self, x1, y1, x2, y2, color: str, width: float = 1.5, dash: bool = False, arrow: bool = False):
        ln = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
        _nostyle(ln)
        ln.line.color.rgb = _rgb(color)
        ln.line.width = Pt(width)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if arrow:
            lnx = ln.line._get_or_add_ln()
            tail = lnx.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
            lnx.append(tail)
        return ln

    def text(self, path: str, x, y, w, h, paras, *, size: float, min_size: float = BODY_MIN, bold: bool = False,
             color: str | None = None, align: str = "left", anchor: str = "top", font: str | None = None,
             line_spacing: float = 1.15, para_space: float = 0.35, role: str = "body", fill: str | None = None,
             fixed: bool = False, italic: bool = False):
        """添加文本框。paras 可以是字符串，或 [{text, level, bullet, bold, k, color}]。"""
        if isinstance(paras, str):
            paras = [{"text": paras}]
        paras = [p if isinstance(p, dict) else {"text": str(p)} for p in paras]
        paras = [p for p in paras if (p.get("text") or "").strip() != "" or p.get("keep")]
        if not paras:
            return None
        family = font or self.t.body_font
        inner_w = (w - 2 * INSET_X) * 72
        inner_h = (h - 2 * INSET_Y) * 72
        # 项目符号占用的缩进
        spec = []
        for p in paras:
            ind = 0.28 * 72 if p.get("bullet") else 0
            ind += 0.3 * 72 * p.get("level", 0)
            spec.append((p["text"], p.get("k", 1.0), ind))
        if fixed:
            fs, over = size, False
            m = measurer(family, bold)
            total = 0.0
            for i, (txt, k, ind) in enumerate(spec):
                total += len(m.wrap(_strip_md(txt), inner_w - ind, fs * k)) * fs * k * line_spacing * m.line_factor + (fs * k * para_space if i else 0)
            over = total > inner_h + 0.5
        else:
            fs, over = _fit(spec, inner_w, inner_h, family, size, min_size, line_spacing, para_space, bold)
        tb = self.slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
        tb.name = f"dw:{self.s.id}:{path}"
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = E(INSET_X)
        tf.margin_top = tf.margin_bottom = E(INSET_Y)
        tf.vertical_anchor = ANCHOR[anchor]
        if fill:
            tb.fill.solid()
            tb.fill.fore_color.rgb = _rgb(fill)
        col = color or self.t.text
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = ALIGN[p.get("align", align)]
            para.line_spacing = line_spacing
            k = p.get("k", 1.0)
            if i:
                para.space_before = Pt(fs * k * para_space)
            if p.get("bullet"):
                _bullet(para, p.get("level", 0), p.get("bullet_color") or self.t.primary, "•" if p.get("level", 0) == 0 else "–")
            self._runs(para, p["text"], fs * k, p.get("bold", bold), p.get("color", col), family, italic=p.get("italic", italic))
        full = "\n".join(_strip_md(p["text"]) for p in paras)
        self.elems.append(Elem(path, self.s.id, "text", self.rel(x, y, w, h), full, fs, min_size, over, role))
        return tb

    def _runs(self, para, text: str, size: float, bold: bool, color: str, family: str, italic: bool = False):
        parts = _MD.split(text)
        for j, seg in enumerate(parts):
            if not seg:
                continue
            r = para.add_run()
            r.text = seg
            f = r.font
            f.size = Pt(size)
            f.bold = bold or (j % 2 == 1)
            f.italic = italic
            f.color.rgb = _rgb(color)
            _set_fonts(r, self.t.latin_font, family)

    def picture(self, path: str, ref: ImageRef | dict | None, x, y, w, h, *, mode: str = "cover"):
        if isinstance(ref, dict):
            ref = ImageRef.model_validate(ref)
        src = self.assets.get(ref.asset) if (ref and ref.asset) else None
        if not src or not Path(src).exists():
            # 占位：浅色底 + 图片图标，提示可替换
            ph = self.rect(x, y, w, h, mix(self.t.surface, self.t.primary, 0.08), None, MSO_SHAPE.RECTANGLE, name=f"dw:{self.s.id}:{path}")
            sz = min(w, h) * 0.3
            add_icon(self.slide.shapes, "search" if not ref else "cloud", E(x + w / 2 - sz / 2), E(y + h / 2 - sz / 2), E(sz), mix(self.t.primary, self.bg, 0.4))
            self.elems.append(Elem(path, self.s.id, "image", self.rel(x, y, w, h), role="image_placeholder"))
            return ph
        img_path = _crop_to(Path(src), w / h if mode == "cover" else None, self.tmp)
        if mode == "contain":
            with Image.open(img_path) as im:
                iw, ih = im.size
            r = min(w / iw, h / ih)
            nw, nh = iw * r, ih * r
            x, y, w, h = x + (w - nw) / 2, y + (h - nh) / 2, nw, nh
        pic = self.slide.shapes.add_picture(str(img_path), E(x), E(y), E(w), E(h))
        pic.name = f"dw:{self.s.id}:{path}"
        self.elems.append(Elem(path, self.s.id, "image", self.rel(x, y, w, h), role="image"))
        return pic

    def chart(self, path: str, spec: ChartSpec | dict, x, y, w, h, font_size: float = 12):
        if isinstance(spec, dict):
            spec = ChartSpec.model_validate(spec)
        label = f"第{self.index + 1}页"
        self.result.all_charts.append((label, spec))
        if spec.native:
            gf = add_native_chart(self.slide, spec, E(x), E(y), E(w), E(h), self.t, font_size)
            gf.name = f"dw:{self.s.id}:{path}"
            self.elems.append(Elem(path, self.s.id, "chart", self.rel(x, y, w, h), role="chart"))
            return gf
        png = chart_png(spec, self.t, width_in=w, height_in=h)
        p = self.tmp / f"chart_{self.s.id}_{len(self.elems)}.png"
        p.write_bytes(png)
        pic = self.slide.shapes.add_picture(str(p), E(x), E(y), E(w), E(h))
        pic.name = f"dw:{self.s.id}:{path}"
        self.result.image_charts.append((label, spec))
        self.elems.append(Elem(path, self.s.id, "image", self.rel(x, y, w, h), role="image_chart"))
        return pic

    def table(self, path: str, spec: TableSpec | dict, x, y, w, h):
        if isinstance(spec, dict):
            spec = TableSpec.model_validate(spec)
        t = self.t
        ncol = len(spec.columns)
        rows = [spec.columns] + [[_cell(v) for v in r] for r in spec.rows]
        nrow = len(rows)
        # 列宽：按内容长度分配
        m = measurer(t.body_font)
        weights = []
        for c in range(ncol):
            lens = [m.width(str(r[c]), 1) for r in rows]
            weights.append(max(2.5, min(18.0, max(lens) if lens else 3)))
        tw = sum(weights)
        colw = [w * k / tw for k in weights]
        # 字号：保证总高度放得下
        fs = 16 if nrow <= 5 else 14 if nrow <= 8 else 12
        fs = min(fs, 16 if ncol <= 4 else 14 if ncol <= 6 else 12)
        while True:
            need = 0.0
            heights = []
            for r in rows:
                lines = max(len(m.wrap(str(r[c]), (colw[c] - 0.16) * 72, fs)) for c in range(ncol))
                rh = (lines * fs * 1.2 * m.line_factor / 1.2 + 10) / 72
                heights.append(rh)
                need += rh
            if need <= h or fs <= 10:
                break
            fs -= 1
        over = need > h + 0.05
        if not over:
            # 行高放宽到最多 0.6 英寸，让表格更易读
            extra = [min(0.6, max(hh, (h / nrow))) for hh in heights]
            if sum(extra) <= h:
                heights = extra
                need = sum(heights)
        gf = self.slide.shapes.add_table(nrow, ncol, E(x), E(y), E(w), E(min(h, need)))
        gf.name = f"dw:{self.s.id}:{path}"
        tbl = gf.table
        # 不使用内置表格样式的条纹，颜色全部显式设置
        tblPr = tbl._tbl.tblPr
        tblPr.set("bandRow", "0")
        tblPr.set("firstRow", "1")
        for c in range(ncol):
            tbl.columns[c].width = E(colw[c])
        head_bg = t.primary
        head_fg = on_color(head_bg)
        stripe = mix(t.background, t.primary, 0.05)
        for r in range(nrow):
            tbl.rows[r].height = E(heights[r])
            for c in range(ncol):
                cell = tbl.cell(r, c)
                cell.margin_left = cell.margin_right = E(0.08)
                cell.margin_top = cell.margin_bottom = E(0.04)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                bg = head_bg if (r == 0 and spec.header) else (stripe if r % 2 == 0 else t.background)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(bg)
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                val = str(rows[r][c])
                num = _is_num(val)
                p.alignment = PP_ALIGN.CENTER if r == 0 else (PP_ALIGN.RIGHT if num else PP_ALIGN.LEFT)
                run = p.add_run()
                run.text = val
                run.font.size = Pt(fs)
                run.font.bold = (r == 0 and spec.header) or (c == 0 and spec.first_col_bold)
                run.font.color.rgb = _rgb(head_fg if (r == 0 and spec.header) else t.text)
                _set_fonts(run, t.latin_font, t.body_font)
        full = "\n".join(" ".join(str(v) for v in r) for r in rows)
        self.elems.append(Elem(path, self.s.id, "table", self.rel(x, y, w, min(h, need)), full, fs, 10, over, "table"))
        return gf

    def icon(self, name: str, x, y, size, color: str | None = None, badge: str | None = None, number: str | None = None):
        add_icon(self.slide.shapes, name, E(x), E(y), E(size), color or self.t.primary, badge, number)

    # ----- 公共部件 -----

    def background(self, color: str):
        self.bg = color
        fill = self.slide.background.fill
        fill.solid()
        fill.fore_color.rgb = _rgb(color)

    def title(self, text: str | None = None, *, y: float = 0.42, h: float = 0.85):
        t = self.t
        text = self.s.title if text is None else text
        x = self.mx
        hc = heading_color(t)
        if t.decor == "bar":
            self.rect(x, y + 0.14, 0.09, h - 0.28, t.primary, name=decor_name(self.s.id))
            x += 0.28
        elif t.decor == "band":
            band = mix(t.background, t.primary, 0.07)
            self.rect(0, 0, self.W, y + h + 0.15, band)
            self.rect(0, y + h + 0.15, self.W, 0.04, t.primary)
        elif t.decor == "corner":
            tri = self.slide.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, 0, 0, E(0.9), E(0.9))
            tri.rotation = 90
            tri.fill.solid()
            tri.fill.fore_color.rgb = _rgb(t.primary)
            tri.line.fill.background()
        self.text("title", x, y, self.W - x - self.mx - (1.2 if self.deck.logo_asset else 0), h, text, size=30, min_size=20,
                  bold=True, color=hc, anchor="middle", font=t.heading_font, role="title")
        if t.decor == "underline":
            self.rect(self.mx, y + h + 0.02, 0.9, 0.06, t.accent)

    def footer(self):
        t = self.t
        col = ensure_contrast(t.muted, self.bg, 3.0)
        self.text("_page", self.W - self.mx - 1.2, self.H - 0.45, 1.2, 0.3, f"{self.index + 1}", size=10, min_size=9,
                  color=col, align="right", role="decor", fixed=True)
        if self.deck.footer:
            self.text("_footer", self.mx, self.H - 0.45, self.W / 2, 0.3, self.deck.footer, size=10, min_size=9,
                      color=col, role="decor", fixed=True)
        if self.deck.logo_asset and self.deck.logo_asset in self.assets:
            self.picture("_logo", ImageRef(asset=self.deck.logo_asset), self.W - self.mx - 1.0, 0.45, 1.0, 0.6, mode="contain")
        for e in self.elems:
            if e.id.startswith("_"):
                e.decor = True

    def bullets_paras(self, bullets, *, bullet=True) -> list[dict]:
        out = []
        for b in bullets:
            if isinstance(b, str):
                out.append({"text": b, "bullet": bullet})
                continue
            out.append({"text": b.get("text", ""), "bullet": bullet})
            for sub in b.get("sub") or []:
                out.append({"text": sub, "bullet": bullet, "level": 1, "k": 0.88, "color": self.t.muted if contrast(self.t.muted, self.bg) >= 4.5 else self.t.text})
        return out


def decor_name(sid: str) -> str:
    return f"decor:{sid}"


# ---------- 布局 ----------

def L_cover(r: SlideRenderer, c: dict):
    t = r.t
    W, H = r.W, r.H
    img = c.get("image")
    has_img = bool(img and img.get("asset") and img.get("asset") in r.assets)
    style = t.cover
    if style == "image" and not has_img:
        style = "solid"
    sub, meta = c.get("subtitle", ""), " · ".join(v for v in (c.get("author"), c.get("date")) if v)
    if style == "solid":
        r.background(t.primary)
        fg = on_color(t.primary)
        r.rect(0, H - 0.35, W, 0.35, mix(t.primary, "#000000", 0.2))
        r.rect(r.mx, 2.3, 1.2, 0.08, t.accent)
        r.text("title", r.mx, 2.5, W - 2 * r.mx, 1.8, r.s.title, size=44, min_size=28, bold=True, color=fg, font=t.heading_font, anchor="top", role="title")
        r.text("content.subtitle", r.mx, 4.35, W - 2 * r.mx, 0.9, sub, size=22, min_size=16, color=mix(fg, t.primary, 0.15))
        r.text("content.meta", r.mx, H - 1.35, W - 2 * r.mx, 0.5, meta, size=16, min_size=12, color=mix(fg, t.primary, 0.25))
    elif style == "split":
        lw = W * 0.56
        r.rect(0, 0, lw, H, t.primary)
        fg = on_color(t.primary)
        if has_img:
            r.picture("content.image", img, lw, 0, W - lw, H)
        else:
            r.rect(lw, 0, W - lw, H, t.surface)
            r.rect(lw + (W - lw) * 0.25, H * 0.2, (W - lw) * 0.7, (W - lw) * 0.7, mix(t.secondary, t.surface, 0.55), shape=MSO_SHAPE.OVAL)
            r.rect(lw + (W - lw) * 0.1, H * 0.55, (W - lw) * 0.35, (W - lw) * 0.35, mix(t.accent, t.surface, 0.3), shape=MSO_SHAPE.OVAL)
        r.rect(r.mx, 2.25, 1.0, 0.08, t.accent)
        r.text("title", r.mx, 2.45, lw - r.mx - 0.4, 2.0, r.s.title, size=40, min_size=26, bold=True, color=fg, font=t.heading_font, role="title")
        r.text("content.subtitle", r.mx, 4.5, lw - r.mx - 0.4, 0.9, sub, size=20, min_size=14, color=mix(fg, t.primary, 0.15))
        r.text("content.meta", r.mx, H - 1.2, lw - r.mx - 0.4, 0.5, meta, size=14, min_size=11, color=mix(fg, t.primary, 0.25))
    elif style == "image":
        r.picture("content.image", img, 0, 0, W, H)
        r.rect(0, 0, W, H, "#000000", alpha=0.5)
        r.bg = "#333333"
        r.rect(r.mx, 2.3, 1.2, 0.08, t.accent)
        r.text("title", r.mx, 2.5, W - 2 * r.mx, 1.8, r.s.title, size=44, min_size=28, bold=True, color="#FFFFFF", font=t.heading_font, role="title")
        r.text("content.subtitle", r.mx, 4.35, W - 2 * r.mx, 0.9, sub, size=22, min_size=16, color="#E5E7EB")
        r.text("content.meta", r.mx, H - 1.3, W - 2 * r.mx, 0.5, meta, size=16, min_size=12, color="#D1D5DB")
    else:  # frame
        r.background(t.background)
        r.rect(0.35, 0.35, W - 0.7, H - 0.7, None, t.primary, line_w=1.5)
        r.rect(W / 2 - 0.6, 2.2, 1.2, 0.07, t.accent)
        r.text("title", 1.0, 2.4, W - 2.0, 1.8, r.s.title, size=42, min_size=26, bold=True, color=heading_color(t), align="center", font=t.heading_font, role="title")
        r.text("content.subtitle", 1.0, 4.3, W - 2.0, 0.8, sub, size=20, min_size=14, color=t.muted, align="center")
        r.text("content.meta", 1.0, H - 1.35, W - 2.0, 0.5, meta, size=14, min_size=11, color=t.muted, align="center")


def L_toc(r: SlideRenderer, c: dict):
    t = r.t
    W, H = r.W, r.H
    lw = W * 0.3
    r.rect(0, 0, lw, H, t.primary)
    fg = on_color(t.primary)
    r.text("title", r.mx * 0.8, 2.6, lw - r.mx, 1.0, r.s.title or "目录", size=36, min_size=24, bold=True, color=fg, font=t.heading_font, role="title")
    r.text("_contents", r.mx * 0.8, 3.55, lw - r.mx, 0.5, "CONTENTS", size=14, min_size=10, color=mix(fg, t.primary, 0.3), role="decor", fixed=True)
    items = c["items"]
    ncol = 2 if len(items) > 5 else 1
    per = (len(items) + ncol - 1) // ncol
    x0 = lw + 0.7
    colw = (W - x0 - r.mx) / ncol
    rowh = min(0.95, (H - 1.6) / per)
    y0 = (H - rowh * per) / 2
    for i, it in enumerate(items):
        col, row = divmod(i, per)
        x = x0 + col * colw
        y = y0 + row * rowh
        r.rect(x, y + rowh * 0.15, 0.62, rowh * 0.7, mix(t.primary, t.background, 0.85), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        r.text(f"content.items.{i}._n", x, y + rowh * 0.15, 0.62, rowh * 0.7, f"{i + 1:02d}", size=18, min_size=12, bold=True, color=t.primary, align="center", anchor="middle", role="decor", fixed=True)
        r.text(f"content.items.{i}", x + 0.8, y, colw - 1.0, rowh, it, size=22, min_size=14, anchor="middle", color=t.text)


def L_section(r: SlideRenderer, c: dict):
    t = r.t
    W, H = r.W, r.H
    r.background(t.primary)
    fg = on_color(t.primary)
    r.rect(W * 0.62, -1.0, W * 0.6, W * 0.6, mix(t.primary, "#FFFFFF", 0.1), shape=MSO_SHAPE.OVAL)
    num = c.get("number") or f"{r.index:02d}"
    r.text("content.number", r.mx, 1.75, 4, 1.55, num, size=72, min_size=40, bold=True, color=ensure_contrast(t.accent, t.primary, 3.0), font=t.heading_font, role="title", fixed=True, line_spacing=1.0)
    r.text("title", r.mx, 3.3, W - 2 * r.mx, 1.2, r.s.title, size=40, min_size=26, bold=True, color=fg, font=t.heading_font, role="title")
    r.text("content.subtitle", r.mx, 4.5, W - 2 * r.mx, 0.9, c.get("subtitle", ""), size=20, min_size=14, color=mix(fg, t.primary, 0.2))


def L_ending(r: SlideRenderer, c: dict):
    t = r.t
    W, H = r.W, r.H
    r.background(t.primary)
    fg = on_color(t.primary)
    r.rect(W / 2 - 0.6, 2.35, 1.2, 0.07, t.accent)
    r.text("title", 1.0, 2.5, W - 2, 1.4, r.s.title or "谢谢", size=48, min_size=28, bold=True, color=fg, align="center", font=t.heading_font, role="title")
    r.text("content.subtitle", 1.0, 3.95, W - 2, 0.8, c.get("subtitle", ""), size=20, min_size=14, color=mix(fg, t.primary, 0.15), align="center")
    r.text("content.contact", 1.0, H - 1.5, W - 2, 0.6, c.get("contact", ""), size=14, min_size=11, color=mix(fg, t.primary, 0.25), align="center")


def L_bullets(r: SlideRenderer, c: dict):
    r.title()
    paras = r.bullets_paras(c["bullets"])
    short = len(paras) <= 3 and all(len(p["text"]) < 30 for p in paras)
    r.text("content.bullets", r.mx + 0.1, r.cy + 0.1, r.cw - 0.2, r.ch - 0.2, paras,
           size=28 if short else 22, min_size=BODY_MIN, para_space=0.6, line_spacing=1.2)


def _column(r: SlideRenderer, key: str, col: dict, x, y, w, h, head_color: str):
    t = r.t
    r.text(f"content.{key}.heading", x, y, w, 0.6, col.get("heading", ""), size=22, min_size=16, bold=True, color=head_color, font=t.heading_font)
    r.rect(x + 0.05, y + 0.65, 0.7, 0.05, t.accent)
    r.text(f"content.{key}.bullets", x, y + 0.85, w, h - 0.85, r.bullets_paras(col.get("bullets", [])), size=20, min_size=BODY_MIN, para_space=0.5)


def L_two_column(r: SlideRenderer, c: dict):
    r.title()
    gap = 0.6
    w = (r.cw - gap) / 2
    _column(r, "left", c["left"], r.mx, r.cy + 0.1, w, r.ch - 0.1, heading_color(r.t))
    r.line(r.mx + w + gap / 2, r.cy + 0.2, r.mx + w + gap / 2, r.cb - 0.2, mix(r.t.muted, r.t.background, 0.6), 1)
    _column(r, "right", c["right"], r.mx + w + gap, r.cy + 0.1, w, r.ch - 0.1, heading_color(r.t))


def L_cards(r: SlideRenderer, c: dict):
    t = r.t
    r.title()
    cards = c["cards"]
    n = len(cards)
    grid = n == 4 and r.W < 11
    cols = 2 if grid else n
    rows = 2 if grid else 1
    gap = 0.35
    w = (r.cw - gap * (cols - 1)) / cols
    h = min((r.ch - gap * (rows - 1)) / rows, 3.9 if rows == 1 else 2.6)
    y0 = r.cy + (r.ch - (h * rows + gap * (rows - 1))) / 2
    card_bg = t.surface
    for i, cd in enumerate(cards):
        row, col = divmod(i, cols)
        x = r.mx + col * (w + gap)
        y = y0 + row * (h + gap)
        r.rect(x, y, w, h, card_bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        r.rect(x, y, w, 0.08, [t.primary, t.secondary, t.accent, t.primary][i % 4])
        isz = min(0.85, h * 0.2)
        r.icon(cd.get("icon") or "", x + 0.3, y + 0.35, isz, t.primary, badge=mix(t.primary, card_bg, 0.85), number=str(i + 1))
        ty = y + 0.35 + isz + 0.2
        r.text(f"content.cards.{i}.heading", x + 0.25, ty, w - 0.5, 0.65, cd["heading"], size=20, min_size=15, bold=True, color=heading_color(t), font=t.heading_font)
        r.text(f"content.cards.{i}.body", x + 0.25, ty + 0.7, w - 0.5, y + h - (ty + 0.7) - 0.2, cd.get("body", ""), size=18, min_size=BODY_MIN, color=t.text, line_spacing=1.25)


def L_quote(r: SlideRenderer, c: dict):
    t = r.t
    r.background(t.surface)
    W, H = r.W, r.H
    r.text("_mark", r.mx, 0.9, 2, 1.8, "“", size=120, min_size=60, bold=True, color=mix(t.accent, t.surface, 0.2), font="Georgia", role="decor", fixed=True)
    r.text("title", r.mx, 0.45, W - 2 * r.mx, 0.6, r.s.title, size=18, min_size=14, color=t.muted, role="title")
    r.text("content.quote", 1.4, 2.0, W - 2.8, 3.0, c["quote"], size=32, min_size=18, bold=True, color=heading_color(t), align="center", anchor="middle", font=t.heading_font, line_spacing=1.3)
    if c.get("source"):
        r.text("content.source", 1.4, 5.2, W - 2.8, 0.6, "—— " + c["source"], size=18, min_size=14, color=t.muted, align="right")


def L_image_text(r: SlideRenderer, c: dict):
    r.title()
    iw = r.cw * 0.46
    gap = 0.45
    left_img = c.get("image_side") == "left"
    ix = r.mx if left_img else r.mx + r.cw - iw
    tx = r.mx + iw + gap if left_img else r.mx
    cap_h = 0.45 if c.get("caption") else 0
    r.picture("content.image", c["image"], ix, r.cy, iw, r.ch - cap_h)
    if cap_h:
        r.text("content.caption", ix, r.cb - cap_h + 0.05, iw, cap_h, c["caption"], size=12, min_size=10, color=r.t.muted, align="center", role="caption")
    r.text("content.bullets", tx, r.cy + 0.1, r.cw - iw - gap, r.ch - 0.2, r.bullets_paras(c.get("bullets", [])), size=22, min_size=BODY_MIN, para_space=0.6, anchor="middle")


def L_full_image(r: SlideRenderer, c: dict):
    W, H = r.W, r.H
    r.picture("content.image", c["image"], 0, 0, W, H)
    bh = 1.5 if c.get("caption") else 1.1
    r.rect(0, H - bh, W, bh, "#000000", alpha=0.55)
    r.bg = "#2B2B2B"
    r.text("title", r.mx, H - bh + 0.15, W - 2 * r.mx, 0.7, r.s.title, size=30, min_size=20, bold=True, color="#FFFFFF", font=r.t.heading_font, role="title")
    if c.get("caption"):
        r.text("content.caption", r.mx, H - bh + 0.85, W - 2 * r.mx, 0.5, c["caption"], size=16, min_size=12, color="#E5E7EB")


def L_image_grid(r: SlideRenderer, c: dict):
    r.title()
    imgs = c["images"]
    n = len(imgs)
    cols = n if n <= 3 else (2 if n == 4 else 3)
    rows = (n + cols - 1) // cols
    gap = 0.3
    w = (r.cw - gap * (cols - 1)) / cols
    h = (r.ch - gap * (rows - 1)) / rows
    for i, it in enumerate(imgs):
        row, col = divmod(i, cols)
        x = r.mx + col * (w + gap)
        y = r.cy + row * (h + gap)
        cap = 0.4 if it.get("caption") else 0
        r.picture(f"content.images.{i}.image", it["image"], x, y, w, h - cap)
        if cap:
            r.text(f"content.images.{i}.caption", x, y + h - cap, w, cap, it["caption"], size=13, min_size=10, color=r.t.muted, align="center", role="caption")


def L_process(r: SlideRenderer, c: dict):
    t = r.t
    r.title()
    steps = c["steps"]
    n = len(steps)
    colw = r.cw / n
    cy = r.cy + 0.9
    d = min(1.0, colw * 0.45)
    r.line(r.mx + colw / 2, cy, r.mx + r.cw - colw / 2, cy, mix(t.primary, t.background, 0.6), 3)
    for i, st in enumerate(steps):
        x = r.mx + i * colw
        cx = x + colw / 2
        fill = t.primary if i % 2 == 0 else t.secondary
        r.rect(cx - d / 2, cy - d / 2, d, d, fill, shape=MSO_SHAPE.OVAL)
        r.text(f"content.steps.{i}._n", cx - d / 2, cy - d / 2, d, d, str(i + 1), size=24, min_size=14, bold=True, color=on_color(fill), align="center", anchor="middle", role="decor", fixed=True)
        if i < n - 1:
            ax = cx + d / 2 + 0.08
            r.line(ax, cy, x + colw * 1.5 - d / 2 - 0.1, cy, t.primary, 2, arrow=True)
        r.text(f"content.steps.{i}.label", x + 0.1, cy + d / 2 + 0.25, colw - 0.2, 0.7, st["label"], size=20, min_size=14, bold=True, color=heading_color(t), align="center", font=t.heading_font)
        r.text(f"content.steps.{i}.detail", x + 0.1, cy + d / 2 + 1.0, colw - 0.2, r.cb - (cy + d / 2 + 1.0), st.get("detail", ""), size=16, min_size=BODY_MIN, color=t.text, align="center", line_spacing=1.25)


def L_timeline(r: SlideRenderer, c: dict):
    t = r.t
    r.title()
    ev = c["events"]
    n = len(ev)
    axis_y = r.cy + r.ch / 2
    r.line(r.mx, axis_y, r.mx + r.cw, axis_y, t.primary, 3)
    colw = r.cw / n
    half = r.ch / 2 - 0.25
    for i, e in enumerate(ev):
        cx = r.mx + colw * (i + 0.5)
        r.rect(cx - 0.14, axis_y - 0.14, 0.28, 0.28, t.accent if i % 2 else t.primary, shape=MSO_SHAPE.OVAL)
        up = i % 2 == 0
        bw = min(colw * 1.8, 3.2) if n > 3 else colw - 0.3
        bx = max(r.mx, min(cx - bw / 2, r.mx + r.cw - bw))
        if up:
            by = r.cy
            r.line(cx, axis_y - 0.14, cx, axis_y - 0.45, mix(t.primary, t.background, 0.5), 1.2)
            r.text(f"content.events.{i}.date", bx, by, bw, 0.5, e["date"], size=18, min_size=13, bold=True, color=t.primary, align="center")
            r.text(f"content.events.{i}.label", bx, by + 0.5, bw, 0.55, e["label"], size=17, min_size=14, bold=True, color=t.text, align="center")
            r.text(f"content.events.{i}.detail", bx, by + 1.05, bw, half - 1.3, e.get("detail", ""), size=14, min_size=12, color=t.muted, align="center")
        else:
            by = axis_y + 0.45
            r.line(cx, axis_y + 0.14, cx, axis_y + 0.45, mix(t.primary, t.background, 0.5), 1.2)
            r.text(f"content.events.{i}.date", bx, by, bw, 0.5, e["date"], size=18, min_size=13, bold=True, color=t.primary, align="center")
            r.text(f"content.events.{i}.label", bx, by + 0.5, bw, 0.55, e["label"], size=17, min_size=14, bold=True, color=t.text, align="center")
            r.text(f"content.events.{i}.detail", bx, by + 1.05, bw, r.cb - by - 1.05, e.get("detail", ""), size=14, min_size=12, color=t.muted, align="center")


def L_matrix(r: SlideRenderer, c: dict):
    t = r.t
    r.title()
    lab = 0.5
    x0, y0 = r.mx + lab, r.cy
    w = r.cw - lab
    h = r.ch - lab
    gap = 0.15
    qw, qh = (w - gap) / 2, (h - gap) / 2
    tints = [mix(t.secondary, t.background, 0.85), mix(t.accent, t.background, 0.7), mix(t.muted, t.background, 0.88), mix(t.primary, t.background, 0.85)]
    for i, q in enumerate(c["quadrants"]):
        row, col = divmod(i, 2)
        x = x0 + col * (qw + gap)
        y = y0 + row * (qh + gap)
        r.rect(x, y, qw, qh, tints[i], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
        r.text(f"content.quadrants.{i}.heading", x + 0.25, y + 0.2, qw - 0.5, 0.6, q["heading"], size=20, min_size=15, bold=True, color=heading_color(t))
        r.text(f"content.quadrants.{i}.body", x + 0.25, y + 0.85, qw - 0.5, qh - 1.0, q.get("body", ""), size=16, min_size=BODY_MIN, color=t.text)
    if c.get("y_label"):
        tb = r.text("content.y_label", r.mx - 1.0 * (h - 0.5) / 2 + 0.25, y0 + h / 2 - 0.2, h, 0.4, "↑ " + c["y_label"], size=14, min_size=11, color=t.muted, align="center", role="caption")
        if tb is not None:
            tb.rotation = -90
            e = r.elems[-1]
            cx, cy = e.box[0] + e.box[2] / 2, e.box[1] + e.box[3] / 2
            w_, h_ = e.box[3] * r.H / r.W, e.box[2] * r.W / r.H
            e.box = [round(cx - w_ / 2, 5), round(cy - h_ / 2, 5), round(w_, 5), round(h_, 5)]
    if c.get("x_label"):
        r.text("content.x_label", x0, r.cb - lab + 0.1, w, 0.4, c["x_label"] + " →", size=14, min_size=11, color=t.muted, align="center", role="caption")


def L_hierarchy(r: SlideRenderer, c: dict):
    t = r.t
    r.title()
    kids = c["children"]
    n = len(kids)
    rw, rh = min(4.0, r.cw * 0.4), 0.8
    rx = r.mx + (r.cw - rw) / 2
    ry = r.cy
    r.rect(rx, ry, rw, rh, t.primary, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    r.text("content.root", rx, ry, rw, rh, c["root"], size=22, min_size=14, bold=True, color=on_color(t.primary), align="center", anchor="middle", font=t.heading_font)
    gap = 0.25
    cw = (r.cw - gap * (n - 1)) / n
    ky = ry + rh + 0.7
    busy = ry + rh + 0.35
    r.line(r.mx + cw / 2, busy, r.mx + r.cw - cw / 2, busy, t.muted, 1.5)
    r.line(r.mx + r.cw / 2, ry + rh, r.mx + r.cw / 2, busy, t.muted, 1.5)
    for i, k in enumerate(kids):
        x = r.mx + i * (cw + gap)
        r.line(x + cw / 2, busy, x + cw / 2, ky, t.muted, 1.5)
        r.rect(x, ky, cw, 0.7, t.secondary, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
        r.text(f"content.children.{i}.label", x, ky, cw, 0.7, k["label"], size=18, min_size=13, bold=True, color=on_color(t.secondary), align="center", anchor="middle")
        if k.get("items"):
            ph = min(r.cb - ky - 0.8, 0.45 + 0.55 * len(k["items"]) * (1.6 if cw < 2.2 else 1.0))
            r.rect(x, ky + 0.8, cw, ph, t.surface, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
            r.text(f"content.children.{i}.items", x + 0.1, ky + 0.9, cw - 0.2, ph - 0.2,
                   [{"text": s, "bullet": True} for s in k["items"]], size=16, min_size=12, color=t.text)


def L_comparison(r: SlideRenderer, c: dict):
    t = r.t
    r.title()
    vh = 0.9 if c.get("verdict") else 0
    gap = 0.4
    w = (r.cw - gap) / 2
    h = r.ch - vh - (0.2 if vh else 0)
    for key, x, colr in (("left", r.mx, t.secondary), ("right", r.mx + w + gap, t.primary)):
        col = c[key]
        r.rect(x, r.cy, w, h, mix(colr, t.background, 0.9), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        r.rect(x, r.cy, w, 0.75, colr, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        r.text(f"content.{key}.heading", x + 0.2, r.cy, w - 0.4, 0.75, col.get("heading", ""), size=22, min_size=15, bold=True, color=on_color(colr), anchor="middle", align="center", font=t.heading_font)
        r.text(f"content.{key}.bullets", x + 0.3, r.cy + 0.95, w - 0.6, h - 1.1, r.bullets_paras(col.get("bullets", [])), size=19, min_size=BODY_MIN, para_space=0.5)
    r.text("_vs", r.mx + w + gap / 2 - 0.35, r.cy + h / 2 - 0.35, 0.7, 0.7, "VS", size=18, min_size=12, bold=True, color=on_color(t.accent), fill=t.accent, align="center", anchor="middle", role="decor", fixed=True)
    if vh:
        r.text("content.verdict", r.mx, r.cb - vh, r.cw, vh, c["verdict"], size=18, min_size=14, bold=True, color=t.text,
               fill=mix(t.accent, t.background, 0.75), anchor="middle", align="center")


def L_chart(r: SlideRenderer, c: dict):
    r.title()
    bl = c.get("bullets") or []
    cw = r.cw * 0.62 if bl else r.cw
    r.chart("content.chart", c["chart"], r.mx, r.cy, cw, r.ch, font_size=13)
    if bl:
        x = r.mx + cw + 0.4
        w = r.cw - cw - 0.4
        r.rect(x, r.cy, w, r.ch, r.t.surface, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
        r.text("content.bullets", x + 0.25, r.cy + 0.25, w - 0.5, r.ch - 0.5, [{"text": b, "bullet": True} for b in bl], size=18, min_size=BODY_MIN, para_space=0.6, anchor="middle")
    src = (c.get("chart") or {}).get("source")
    if src:
        r.text("content.chart.source", r.mx, r.cb, r.cw, 0.3, "数据来源：" + src, size=10, min_size=9, color=r.t.muted, role="caption", fixed=True)


def L_big_number(r: SlideRenderer, c: dict):
    t = r.t
    r.title()
    ms = c["metrics"]
    n = len(ms)
    w = r.cw / n
    top = r.cy + 0.5
    for i, m in enumerate(ms):
        x = r.mx + i * w
        if i:
            r.line(x, top + 0.2, x, r.cb - 0.6, mix(t.muted, t.background, 0.6), 1)
        r.text(f"content.metrics.{i}.value", x + 0.15, top, w - 0.3, 1.5, m["value"], size=60 if n <= 2 else 48, min_size=28, bold=True, color=heading_color(t), align="center", anchor="bottom", font=t.heading_font)
        r.text(f"content.metrics.{i}.label", x + 0.15, top + 1.6, w - 0.3, 0.6, m["label"], size=20, min_size=14, bold=True, color=t.text, align="center")
        r.text(f"content.metrics.{i}.detail", x + 0.25, top + 2.25, w - 0.5, r.cb - top - 2.3, m.get("detail", ""), size=15, min_size=12, color=t.muted, align="center")


def L_table(r: SlideRenderer, c: dict):
    r.title()
    cap = (c["table"] or {}).get("caption")
    ch = r.ch - (0.4 if cap else 0)
    r.table("content.table", c["table"], r.mx, r.cy, r.cw, ch)
    if cap:
        r.text("content.table.caption", r.mx, r.cb - 0.35, r.cw, 0.35, cap, size=12, min_size=10, color=r.t.muted, role="caption")


def L_chart_table(r: SlideRenderer, c: dict):
    r.title()
    cw = r.cw * 0.55
    r.chart("content.chart", c["chart"], r.mx, r.cy, cw, r.ch, font_size=12)
    r.table("content.table", c["table"], r.mx + cw + 0.35, r.cy + 0.2, r.cw - cw - 0.35, r.ch - 0.4)


def L_team(r: SlideRenderer, c: dict):
    t = r.t
    r.title()
    ps = c["people"]
    n = len(ps)
    cols = n if n <= 4 else 3
    rows = (n + cols - 1) // cols
    gap = 0.3
    w = (r.cw - gap * (cols - 1)) / cols
    h = min((r.ch - gap * (rows - 1)) / rows, 3.6)
    y0 = r.cy + (r.ch - (h * rows + gap * (rows - 1))) / 2
    for i, p in enumerate(ps):
        row, col = divmod(i, cols)
        x = r.mx + col * (w + gap)
        y = y0 + row * (h + gap)
        r.rect(x, y, w, h, t.surface, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
        d = min(1.3, h * 0.38, w * 0.5)
        img = p.get("image")
        if img and img.get("asset") and img["asset"] in r.assets:
            r.picture(f"content.people.{i}.image", img, x + (w - d) / 2, y + 0.25, d, d)
        else:
            r.rect(x + (w - d) / 2, y + 0.25, d, d, mix(t.primary, t.surface, 0.2), shape=MSO_SHAPE.OVAL)
            r.text(f"content.people.{i}._initial", x + (w - d) / 2, y + 0.25, d, d, (p["name"] or "?")[:1], size=28, min_size=14, bold=True, color=on_color(t.primary), align="center", anchor="middle", role="decor", fixed=True)
        ty = y + 0.35 + d
        r.text(f"content.people.{i}.name", x + 0.1, ty, w - 0.2, 0.5, p["name"], size=20, min_size=14, bold=True, color=t.text, align="center")
        r.text(f"content.people.{i}.role", x + 0.1, ty + 0.5, w - 0.2, 0.45, p.get("role", ""), size=15, min_size=11, color=t.primary, align="center")
        r.text(f"content.people.{i}.bio", x + 0.2, ty + 1.0, w - 0.4, y + h - ty - 1.1, p.get("bio", ""), size=15, min_size=11, color=t.muted, align="center")


LAYOUTS = {
    "cover": L_cover, "toc": L_toc, "section": L_section, "ending": L_ending, "bullets": L_bullets,
    "two_column": L_two_column, "cards": L_cards, "quote": L_quote, "image_text": L_image_text,
    "full_image": L_full_image, "image_grid": L_image_grid, "process": L_process, "timeline": L_timeline,
    "matrix": L_matrix, "hierarchy": L_hierarchy, "comparison": L_comparison, "chart": L_chart,
    "big_number": L_big_number, "table": L_table, "chart_table": L_chart_table, "team": L_team,
}
NO_FOOTER = {"cover", "section", "ending", "full_image", "toc"}


def render_deck(deck: Deck, out: Path, assets: dict[str, Path] | None = None, tmp: Path | None = None,
                only: list[str] | None = None) -> RenderResult:
    """渲染整份或部分页面（only=页面 ID 列表，用于只重渲变化的预览页）。"""
    assets = {k: Path(v) for k, v in (assets or {}).items()}
    tmp = tmp or out.parent
    tmp.mkdir(parents=True, exist_ok=True)
    deck = deck.model_copy(deep=True)
    deck.theme = normalize_theme(deck.theme)
    prs = Presentation()
    W = 13.333 if deck.aspect == "16:9" else 10.0
    H = 7.5
    prs.slide_width = E(W)
    prs.slide_height = E(H)
    prs.core_properties.title = deck.title or (deck.slides[0].title if deck.slides else "")
    prs.core_properties.author = "DocWork"
    blank = prs.slide_layouts[6]
    res = RenderResult(path=out)
    total = len(deck.slides)
    for idx, s in enumerate(deck.slides):
        if only is not None and s.id not in only:
            continue
        sl = prs.slides.add_slide(blank)
        r = SlideRenderer(deck, sl, s, W, H, idx, total, assets, tmp, res)
        r.background(deck.theme.background)
        LAYOUTS[s.layout](r, s.content)
        if s.layout not in NO_FOOTER:
            r.footer()
        if s.notes:
            sl.notes_slide.notes_text_frame.text = s.notes
        res.slide_ids.append(s.id)
        res.elements[s.id] = [e.__dict__ for e in r.elems]
        for e in r.elems:
            if e.overflow and not e.decor:
                res.issues.append({"slide": s.id, "index": idx, "element": e.id, "kind": "overflow_estimate",
                                   "message": f"文字可能放不下（已缩到 {e.size:.0f}pt）"})
    prs.save(out)
    return res


# ---------- 辅助 ----------

def _fit(spec, inner_w, inner_h, family, size, min_size, line_spacing, para_space, bold):
    m = measurer(family, bold)
    fs = size
    while True:
        total = 0.0
        for i, (txt, k, ind) in enumerate(spec):
            s = fs * k
            total += len(m.wrap(_strip_md(txt), inner_w - ind, s)) * s * line_spacing * m.line_factor
            if i:
                total += s * para_space
        if total <= inner_h or fs <= min_size:
            return fs, total > inner_h + 0.5
        fs = max(min_size, fs - 1)


def _strip_md(s: str) -> str:
    return _MD.sub(r"\1", s or "")


def _set_fonts(run, latin: str, ea: str) -> None:
    rPr = run._r.get_or_add_rPr()
    for tag, face in (("a:latin", latin), ("a:ea", ea), ("a:cs", latin)):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            # a:latin/a:ea/a:cs 必须位于 solidFill 等元素之后的正确顺序：直接追加即可满足 schema 顺序
            rPr.append(el)
        el.set("typeface", face)
    # schema 要求 solidFill 在 latin 之前：重新排序
    fill = rPr.find(qn("a:solidFill"))
    if fill is not None:
        rPr.remove(fill)
        rPr.insert(0, fill)


def _bullet(para, level: int, color: str, char: str) -> None:
    pPr = para._p.get_or_add_pPr()
    indent = 0.28 + 0.3 * level
    pPr.set("marL", str(int(indent * EMU_IN)))
    pPr.set("indent", str(int(-0.28 * EMU_IN)))
    for tag in ("a:buNone", "a:buChar", "a:buClr", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buClr = pPr.makeelement(qn("a:buClr"), {})
    clr = buClr.makeelement(qn("a:srgbClr"), {"val": color.lstrip("#").upper()})
    buClr.append(clr)
    buFont = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    buChar = pPr.makeelement(qn("a:buChar"), {"char": char})
    # 顺序：lnSpc, spcBef, spcAft, buClr..., buFont, buChar
    pPr.append(buClr)
    pPr.append(buFont)
    pPr.append(buChar)
    _order_ppr(pPr)


_PPR_ORDER = ["a:lnSpc", "a:spcBef", "a:spcAft", "a:buClrTx", "a:buClr", "a:buSzTx", "a:buSzPct", "a:buSzPts",
              "a:buFontTx", "a:buFont", "a:buNone", "a:buAutoNum", "a:buChar", "a:buBlip", "a:tabLst", "a:defRPr", "a:extLst"]


def _order_ppr(pPr) -> None:
    rank = {qn(t): i for i, t in enumerate(_PPR_ORDER)}
    kids = list(pPr)
    kids.sort(key=lambda e: rank.get(e.tag, 99))
    for k in kids:
        pPr.remove(k)
    for k in kids:
        pPr.append(k)


def _nostyle(shape) -> None:
    """去掉形状的主题样式引用（避免 LibreOffice/WPS 套用主题阴影等效果），颜色全部显式设置。"""
    el = shape._element.find(qn("p:style"))
    if el is not None:
        shape._element.remove(el)


def _set_alpha(shape, alpha: float) -> None:
    sf = shape.fill._xPr.find(qn("a:solidFill"))
    if sf is None:
        return
    clr = sf[0]
    a = clr.makeelement(qn("a:alpha"), {"val": str(int(alpha * 100000))})
    clr.append(a)


def _crop_to(src: Path, ratio: float | None, tmp: Path) -> Path:
    """按目标比例居中裁切并限制分辨率，输出 JPEG/PNG。"""
    with Image.open(src) as im:
        im.load()
        has_alpha = im.mode in ("RGBA", "LA", "P")
        im = im.convert("RGBA" if has_alpha else "RGB")
        if ratio:
            w, h = im.size
            cur = w / h
            if cur > ratio:
                nw = int(h * ratio)
                im = im.crop(((w - nw) // 2, 0, (w - nw) // 2 + nw, h))
            elif cur < ratio:
                nh = int(w / ratio)
                im = im.crop((0, (h - nh) // 2, w, (h - nh) // 2 + nh))
        im.thumbnail((2000, 2000))
        out = tmp / f"img_{src.stem}_{int((ratio or 0) * 1000)}.{'png' if has_alpha else 'jpg'}"
        if has_alpha:
            im.save(out, "PNG")
        else:
            im.save(out, "JPEG", quality=88)
    return out


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:,.2f}".rstrip("0").rstrip(".") if abs(v) < 1e15 else str(v)
    return str(v)


def _is_num(s: str) -> bool:
    return bool(re.fullmatch(r"[-+]?[\d,]*\.?\d+%?", s.strip())) if s else False
