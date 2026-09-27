"""资料抽取：把上传文件转成带结构的文字（Markdown），并收集图片素材。用于生成时的资料和重建。"""
from __future__ import annotations

import io
from pathlib import Path

from .. import storage


def extract(path: Path, kind: str, workdir: Path, cancel=None, progress=None, with_images: bool = True) -> dict:
    """返回 {text, images: [{path, caption}], report}。"""
    workdir.mkdir(parents=True, exist_ok=True)
    images: list[dict] = []
    report: dict = {}
    if kind == "pdf":
        from . import pdfdocx
        pages, rep = pdfdocx.page_blocks(path, cancel=cancel, progress=progress, image_dir=(workdir / "img") if with_images else None)
        text = pdfdocx.blocks_markdown(pages)
        for pi, blocks in enumerate(pages, 1):
            for b in blocks:
                if b["type"] == "image":
                    images.append({"path": b["path"], "caption": f"第 {pi} 页图片"})
        report = rep
    elif kind in ("docx", "docm"):
        text, images = _docx(path, workdir, with_images)
    elif kind in ("pptx", "pptm"):
        text, images = _pptx(path, workdir, with_images)
    elif kind in ("xlsx", "xlsm"):
        text = _xlsx(path)
    elif kind in ("doc", "odt", "rtf", "ppt", "odp", "xls", "ods"):
        from . import office
        norm = office.normalize_to_ooxml(path, workdir / "norm", cancel=cancel)
        return extract(norm, norm.suffix.lstrip(".").lower(), workdir, cancel, progress, with_images)
    elif kind == "csv":
        text = _csv(path)
    elif kind in ("md", "txt"):
        text = _read_text(path)
    elif kind == "html":
        import re
        raw = _read_text(path)
        text = re.sub(r"<script.*?</script>|<style.*?</style>", "", raw, flags=re.S | re.I)
        text = re.sub(r"<[^>]+>", " ", text)
        text = re.sub(r"[ \t]+", " ", text)
    elif kind in ("png", "jpg", "gif", "bmp", "tiff", "webp"):
        from . import ocr
        lines = ocr.ocr_image(path)
        from PIL import Image
        with Image.open(path) as im:
            w = im.size[0]
        text = ocr.blocks_to_markdown(ocr.analyze(lines, w))
        images.append({"path": str(path), "caption": path.stem})
        report = {"ocr": True, "engine": ocr.engine_name()}
    else:
        text = ""
    return {"text": text[:400000], "images": images[:60], "report": report}


def _read_text(p: Path) -> str:
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _docx(path: Path, workdir: Path, with_images: bool) -> tuple[str, list]:
    from docx import Document

    d = Document(path)
    out = []
    body = d.element.body
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(child, d)
            t = p.text.strip()
            if not t:
                continue
            st = (p.style.name or "") if p.style is not None else ""
            if st.startswith("Heading") or st.startswith("标题"):
                lvl = "".join(ch for ch in st if ch.isdigit()) or "1"
                out.append("#" * min(3, int(lvl)) + " " + t)
            elif st == "Title":
                out.append("# " + t)
            else:
                out.append(t)
        elif tag == "tbl":
            tb = Table(child, d)
            rows = []
            for r in tb.rows:
                cells = []
                for c in r.cells:
                    cells.append(c.text.strip().replace("\n", " ").replace("|", "/"))
                rows.append(cells)
            if rows:
                n = max(len(r) for r in rows)
                out.append("| " + " | ".join(rows[0]) + " |\n|" + "---|" * n + "\n" + "\n".join("| " + " | ".join(r) + " |" for r in rows[1:]))
    images = []
    if with_images:
        (workdir / "img").mkdir(parents=True, exist_ok=True)
        for i, rel in enumerate(r for r in d.part.rels.values() if "image" in r.reltype):
            try:
                blob = rel.target_part.blob
                ext = Path(rel.target_part.partname).suffix or ".png"
                p = workdir / "img" / f"docx_img{i + 1}{ext}"
                p.write_bytes(blob)
                if _usable_image(p):
                    images.append({"path": str(p), "caption": f"文档图片 {i + 1}"})
            except Exception:
                continue
    return "\n\n".join(out), images


def _pptx(path: Path, workdir: Path, with_images: bool) -> tuple[str, list]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(path)
    out = []
    images = []
    (workdir / "img").mkdir(parents=True, exist_ok=True)

    def walk(shapes, acc, si):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes, acc, si)
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                acc.append(sh.text_frame.text.strip())
            if getattr(sh, "has_table", False) and sh.has_table:
                rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in sh.table.rows]
                n = max(len(r) for r in rows)
                acc.append("| " + " | ".join(rows[0]) + " |\n|" + "---|" * n + "\n" + "\n".join("| " + " | ".join(r) + " |" for r in rows[1:]))
            if getattr(sh, "has_chart", False) and sh.has_chart:
                ch = sh.chart
                try:
                    cats = list(ch.plots[0].categories)
                    ser = [(s.name, list(s.values)) for s in ch.plots[0].series]
                    acc.append("图表数据：分类 " + "、".join(map(str, cats)) + "；" + "；".join(f"{n}: {', '.join(map(str, v))}" for n, v in ser))
                except Exception:
                    pass
            if with_images and sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    p = workdir / "img" / f"slide{si}_{sh.shape_id}.{sh.image.ext}"
                    p.write_bytes(sh.image.blob)
                    if _usable_image(p):
                        images.append({"path": str(p), "caption": f"第 {si} 页图片"})
                except Exception:
                    pass

    for i, slide in enumerate(prs.slides, 1):
        acc: list[str] = []
        walk(slide.shapes, acc, i)
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        out.append(f"## 第 {i} 页\n\n" + "\n\n".join(acc) + (f"\n\n（备注：{notes}）" if notes else ""))
    return "\n\n".join(out), images


def _xlsx(path: Path, max_rows: int = 200) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            if len(rows) >= max_rows:
                break
            if r is None or all(v is None for v in r):
                continue
            rows.append(["" if v is None else str(v) for v in r])
        if rows:
            n = max(len(r) for r in rows)
            rows = [r + [""] * (n - len(r)) for r in rows]
            out.append(f"## 工作表：{ws.title}\n\n| " + " | ".join(rows[0]) + " |\n|" + "---|" * n + "\n" +
                       "\n".join("| " + " | ".join(r) + " |" for r in rows[1:]))
    return "\n\n".join(out)


def _csv(path: Path) -> str:
    import pandas as pd

    for enc in ("utf-8-sig", "gb18030"):
        try:
            df = pd.read_csv(path, encoding=enc, nrows=200)
            break
        except UnicodeDecodeError:
            continue
    else:
        return ""
    return df.to_markdown(index=False) if hasattr(df, "to_markdown") and _has_tabulate() else df.to_csv(index=False)


def _has_tabulate() -> bool:
    try:
        import tabulate  # noqa: F401
        return True
    except ImportError:
        return False


def _usable_image(p: Path) -> bool:
    try:
        from PIL import Image
        with Image.open(p) as im:
            w, h = im.size
        return w >= 200 and h >= 150
    except Exception:
        return False
