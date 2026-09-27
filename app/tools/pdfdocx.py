"""PDF → 可编辑内容的版面解析，以及内置的 PDF → DOCX 重建后端。

文字型页面用 pdfplumber 读取带坐标、字号的文字和表格；扫描页用 OCR。
输出统一的块结构，供 DOCX 重建、PPT 重建（交给模型）和 XLSX 表格提取共用。
pdf2docx（已停止积极维护）作为可替换的首选后端，由 convert 模块选择。
"""
from __future__ import annotations

import re
import statistics
from pathlib import Path

from . import ocr as ocrmod
from .pdf import render_page


def page_blocks(pdf_path: Path, cancel=None, progress=None, image_dir: Path | None = None) -> tuple[list[list[dict]], dict]:
    """返回每页的块列表和报告。块：{type: heading|text|table|image, text|rows|path, level?, box}。"""
    import pdfplumber

    report = {"pages": 0, "ocr_pages": [], "tables": 0, "images": 0, "low_confidence": [], "notes": []}
    pages_out: list[list[dict]] = []
    with pdfplumber.open(pdf_path) as pdf:
        n = len(pdf.pages)
        report["pages"] = n
        # 全文正文字号（用于判断标题）
        sizes = []
        for p in pdf.pages[: min(n, 30)]:
            sizes += [round(c["size"], 1) for c in p.chars[:3000]]
        body = statistics.median(sizes) if sizes else 11
        for i, page in enumerate(pdf.pages):
            if cancel and cancel():
                from .sandbox import ToolCancelled
                raise ToolCancelled()
            if progress:
                progress(i / max(1, n), f"解析第 {i + 1}/{n} 页")
            if len(page.chars) > 5:
                blocks = _text_page(page, body, report)
            else:
                img = render_page(pdf_path, i, dpi=250)
                lines = ocrmod.ocr_image(img)
                report["ocr_pages"].append(i + 1)
                report["low_confidence"] += [{"page": i + 1, "text": l["text"], "conf": round(l["conf"], 2)} for l in lines if l["conf"] < 0.8][:20]
                blocks = ocrmod.analyze(lines, img.size[0])
                if lines:
                    med = statistics.median([l["box"][3] - l["box"][1] for l in lines])
                    for b in blocks:
                        if b["type"] == "text" and b.get("size", 0) > med * 1.4 and len(b["text"]) < 60:
                            b["type"] = "heading"
                            b["level"] = 1 if b["size"] > med * 1.8 else 2
                report["tables"] += sum(1 for b in blocks if b["type"] == "table")
            if image_dir is not None:
                blocks += _images(pdf_path, page, i, image_dir, report)
                blocks.sort(key=lambda b: (b.get("order", 0), b["box"][1]))
            pages_out.append(blocks)
    if report["ocr_pages"]:
        report["notes"].append(f"第 {', '.join(map(str, report['ocr_pages'][:20]))} 页为扫描页，经 OCR 识别，可能存在错字")
    return pages_out, report


def _text_page(page, body: float, report: dict) -> list[dict]:
    W = float(page.width)
    tables = []
    try:
        for t in page.find_tables():
            rows = [[(c or "").replace("\n", " ").strip() for c in r] for r in t.extract()]
            rows = [r for r in rows if any(r)]
            if len(rows) >= 2 and max(len(r) for r in rows) >= 2:
                tables.append({"box": list(t.bbox), "rows": rows})
    except Exception:
        pass
    report["tables"] += len(tables)

    def in_table(w):
        cx, cy = (w["x0"] + w["x1"]) / 2, (w["top"] + w["bottom"]) / 2
        return any(t["box"][0] - 2 <= cx <= t["box"][2] + 2 and t["box"][1] - 2 <= cy <= t["box"][3] + 2 for t in tables)

    words = [w for w in page.extract_words(extra_attrs=["size", "fontname"], keep_blank_chars=False, use_text_flow=False) if not in_table(w)]
    # 组成行：top 相近的词归为一行，行内间距过大时拆分（多栏）
    words.sort(key=lambda w: (round(w["top"]), w["x0"]))
    lines: list[dict] = []
    for w in words:
        if lines and abs(lines[-1]["top"] - w["top"]) < max(2.0, w["size"] * 0.4) and w["x0"] - lines[-1]["x1"] < w["size"] * 2.0:
            ln = lines[-1]
            sep = "" if (ocrmod._cjk(ln["text"][-1]) and ocrmod._cjk(w["text"][0])) else " "
            ln["text"] += sep + w["text"]
            ln["x1"] = max(ln["x1"], w["x1"])
            ln["bottom"] = max(ln["bottom"], w["bottom"])
            ln["sizes"].append(w["size"])
            ln["bold"] = ln["bold"] and "Bold" in w["fontname"]
        else:
            lines.append({"text": w["text"], "x0": w["x0"], "x1": w["x1"], "top": w["top"], "bottom": w["bottom"],
                          "sizes": [w["size"]], "bold": "Bold" in w["fontname"]})
    items = []
    for ln in lines:
        size = statistics.median(ln["sizes"])
        items.append({"box": [ln["x0"], ln["top"], ln["x1"], ln["bottom"]], "text": ln["text"], "conf": 1.0,
                      "size": size, "bold": ln["bold"]})
    for t in tables:
        items.append({"box": t["box"], "text": "", "table": t["rows"], "full": True, "conf": 1.0})
    ordered = ocrmod.order_lines(items, W)
    blocks: list[dict] = []
    buf: list[dict] = []

    def flush():
        if not buf:
            return
        txt = ocrmod.to_text(buf)
        size = statistics.median([b.get("size", body) for b in buf])
        bold = all(b.get("bold") for b in buf)
        box = [min(b["box"][0] for b in buf), min(b["box"][1] for b in buf), max(b["box"][2] for b in buf), max(b["box"][3] for b in buf)]
        for para in [p for p in txt.split("\n\n") if p.strip()]:
            kind, level = "text", 0
            if len(para) < 80 and (size >= body * 1.25 or (bold and size >= body * 1.05 and len(para) < 40)):
                kind = "heading"
                level = 1 if size >= body * 1.6 else 2 if size >= body * 1.25 else 3
            blocks.append({"type": kind, "text": para.strip(), "level": level, "box": box, "order": len(blocks)})
        buf.clear()

    prev = None
    for it in ordered:
        if "table" in it:
            flush()
            blocks.append({"type": "table", "rows": it["table"], "box": it["box"], "order": len(blocks)})
            prev = None
            continue
        # 字号或粗细变化时分段（标题与正文分开）
        if prev is not None and (abs(it.get("size", body) - prev.get("size", body)) > 0.8 or it.get("bold") != prev.get("bold")):
            flush()
        buf.append(it)
        prev = it
    flush()
    return blocks


def _images(pdf_path: Path, page, index: int, image_dir: Path, report: dict) -> list[dict]:
    out = []
    imgs = [im for im in page.images if (im["x1"] - im["x0"]) > 40 and (im["bottom"] - im["top"]) > 40]
    if not imgs:
        return out
    image_dir.mkdir(parents=True, exist_ok=True)
    full = render_page(pdf_path, index, dpi=150)
    k = 150 / 72
    for j, im in enumerate(imgs[:20]):
        box = (int(im["x0"] * k), int(im["top"] * k), int(im["x1"] * k), int(im["bottom"] * k))
        crop = full.crop(box)
        p = image_dir / f"p{index + 1:03d}_img{j + 1}.png"
        crop.save(p)
        out.append({"type": "image", "path": str(p), "box": [im["x0"], im["top"], im["x1"], im["bottom"]], "order": 0,
                    "width_ratio": (im["x1"] - im["x0"]) / float(page.width)})
        report["images"] += 1
    return out


def blocks_markdown(pages: list[list[dict]], with_page_marks: bool = True) -> str:
    out = []
    for i, blocks in enumerate(pages, 1):
        if with_page_marks:
            out.append(f"<!-- 第 {i} 页 -->")
        for b in blocks:
            if b["type"] == "heading":
                out.append("#" * max(1, b.get("level", 1)) + " " + b["text"])
            elif b["type"] == "text":
                out.append(b["text"])
            elif b["type"] == "table":
                out.append(ocrmod.blocks_to_markdown([b]))
            elif b["type"] == "image":
                out.append(f"![图片]({Path(b['path']).name})")
    return "\n\n".join(out)


def build_docx(pages: list[list[dict]], out: Path, title: str = "") -> Path:
    """内置后端：把块结构写成 DOCX（段落、标题、表格、图片），按阅读顺序连续排版。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm, Pt

    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    for m in ("left_margin", "right_margin"):
        setattr(sec, m, Mm(25))
    normal = d.styles["Normal"]
    normal.font.size = Pt(11)
    rpr = normal.element.get_or_add_rPr()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:eastAsia"), "宋体")
    tw = sec.page_width - sec.left_margin - sec.right_margin
    for bi, blocks in enumerate(pages):
        for b in blocks:
            if b["type"] == "heading":
                d.add_heading(b["text"], level=min(3, max(1, b.get("level", 1))))
            elif b["type"] == "text":
                d.add_paragraph(re.sub(r"\s*\n\s*", "", b["text"]) if _mostly_cjk(b["text"]) else b["text"].replace("\n", " "))
            elif b["type"] == "table":
                rows = b["rows"]
                ncol = max(len(r) for r in rows)
                t = d.add_table(rows=len(rows), cols=ncol)
                t.style = d.styles["Table Grid"]
                for ri, r in enumerate(rows):
                    for ci in range(ncol):
                        t.cell(ri, ci).text = r[ci] if ci < len(r) else ""
                        if ri == 0:
                            for run in t.cell(ri, ci).paragraphs[0].runs:
                                run.bold = True
                d.add_paragraph()
            elif b["type"] == "image":
                p = d.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                try:
                    p.add_run().add_picture(b["path"], width=int(tw * min(1.0, max(0.2, b.get("width_ratio", 0.6)))))
                except Exception:
                    pass
    if title:
        d.core_properties.title = title
    d.save(out)
    return out


def _mostly_cjk(s: str) -> bool:
    if not s:
        return False
    return sum(1 for c in s if ord(c) > 0x2E80) > len(s) * 0.3
