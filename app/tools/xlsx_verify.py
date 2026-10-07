"""Excel 验证：生成原件保留不动，在副本上用 LibreOffice 重算后检查。

检查项：错误值；公式引用区域是否在规格声明的数据区域内；简单汇总公式与独立计算结果比对；
规格中声明的关键结果（pandas 计算）比对；图表、条件格式、数据验证是否都存在。
同时输出重算后的值，用于网页表格预览。
"""
from __future__ import annotations

import math
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from ..spec.workbook import Workbook
from . import office, ooxml

ERRORS = ("#REF!", "#NAME?", "#DIV/0!", "#VALUE!", "#N/A", "#NUM!", "#NULL!", "Err:")
_REF = re.compile(r"(?:(?:'([^']+)'|([A-Za-z0-9_一-鿿]+))!)?\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
_SIMPLE = re.compile(r"^=\s*(SUM|AVERAGE|MIN|MAX|COUNT|COUNTA)\(\s*\$?([A-Z]{1,3})\$?(\d+)\s*:\s*\$?([A-Z]{1,3})\$?(\d+)\s*\)\s*$", re.I)


def verify(original: Path, spec: Workbook, workdir: Path, expected_counts: dict | None = None, cancel=None) -> dict:
    recalced = office.recalc_copy(original, workdir, cancel=cancel)
    vals = load_workbook(recalced, data_only=True)
    forms = load_workbook(original)
    problems: list[dict] = []
    preview: dict = {}
    declared = {s.name: s.data_range for s in spec.sheets}
    for s in spec.sheets:
        if s.name not in vals.sheetnames:
            problems.append({"sheet": s.name, "kind": "missing_sheet", "message": f"工作表“{s.name}”不存在"})
            continue
        wv, wf = vals[s.name], forms[s.name]
        nrows, ncols = declared[s.name]
        grid = []
        for r in range(1, min(nrows, 500) + 1):
            row = []
            for c in range(1, min(ncols, 50) + 1):
                v = wv.cell(row=r, column=c).value
                row.append(_jsonable(v))
            grid.append(row)
        preview[s.name] = {"id": s.id, "rows": grid, "truncated": nrows > 500 or ncols > 50,
                           "formats": [c.number_format for c in s.columns]}
        for row in wf.iter_rows(min_row=1, max_row=nrows, max_col=ncols):
            for cell in row:
                f = cell.value
                v = wv[cell.coordinate].value
                if isinstance(v, str) and v.startswith(ERRORS):
                    problems.append({"sheet": s.name, "cell": cell.coordinate, "kind": "error_value", "message": f"{cell.coordinate} 计算结果为 {v}"})
                if not (isinstance(f, str) and f.startswith("=")):
                    continue
                for m in _REF.finditer(f.split('"')[0] if f.count('"') % 2 else re.sub(r'"[^"]*"', '""', f)):
                    sheet_name = m.group(1) or m.group(2) or s.name
                    if sheet_name not in declared:
                        if m.group(2) and not (m.group(1) or "!" in f):
                            continue
                        problems.append({"sheet": s.name, "cell": cell.coordinate, "kind": "bad_ref", "message": f"{cell.coordinate} 引用了不存在的工作表 {sheet_name}"})
                        continue
                    rr, cc = declared[sheet_name]
                    c1, r1 = column_index_from_string(m.group(3)), int(m.group(4))
                    c2 = column_index_from_string(m.group(5)) if m.group(5) else c1
                    r2 = int(m.group(6)) if m.group(6) else r1
                    if max(r1, r2) > rr or max(c1, c2) > cc:
                        problems.append({"sheet": s.name, "cell": cell.coordinate, "kind": "ref_out_of_range",
                                         "message": f"{cell.coordinate} 的公式引用 {m.group(0)} 超出了数据区域（{get_column_letter(cc)}{rr}）"})
                sm = _SIMPLE.match(f)
                if sm and isinstance(v, (int, float)):
                    func = sm.group(1).upper()
                    a1, b1, a2, b2 = sm.group(2), int(sm.group(3)), sm.group(4), int(sm.group(5))
                    nums, nonempty = [], 0
                    for rr_ in range(b1, b2 + 1):
                        for cc_ in range(column_index_from_string(a1), column_index_from_string(a2) + 1):
                            x = wv.cell(row=rr_, column=cc_).value
                            if x not in (None, ""):
                                nonempty += 1
                            if isinstance(x, (int, float)) and not isinstance(x, bool):
                                nums.append(float(x))
                    calc = {"SUM": sum(nums), "AVERAGE": (sum(nums) / len(nums)) if nums else None, "MIN": min(nums) if nums else 0,
                            "MAX": max(nums) if nums else 0, "COUNT": len(nums), "COUNTA": nonempty}[func]
                    if calc is not None and not _close(calc, v):
                        problems.append({"sheet": s.name, "cell": cell.coordinate, "kind": "mismatch",
                                         "message": f"{cell.coordinate} 的 {func} 结果 {v} 与独立计算 {calc:g} 不一致"})
        for e in s.expected:
            v = wv[e.cell].value
            if not isinstance(v, (int, float)) or not _close(e.value, v, e.tolerance):
                problems.append({"sheet": s.name, "cell": e.cell, "kind": "expected_mismatch",
                                 "message": f"{e.cell} 的结果 {v} 与独立计算的 {e.value:g} 不一致"})
    inv = ooxml.inventory(original)
    if expected_counts:
        for k, label, invk in (("charts", "图表", "charts"), ("cond_formats", "条件格式", "cond_formats"), ("validations", "数据验证", "data_validations")):
            if expected_counts.get(k, 0) > inv.get(invk, 0):
                problems.append({"kind": "missing_element", "message": f"{label}应有 {expected_counts[k]} 个，文件中只有 {inv.get(invk, 0)} 个"})
    return {"ok": not problems, "problems": problems, "preview": preview, "inventory": inv}


def _close(a: float, b: float, tol: float = 1e-6) -> bool:
    return math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=max(tol, 1e-9))


def _jsonable(v):
    if v is None or isinstance(v, (int, float, str, bool)):
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return None
        return v
    return str(v)
