"""2026-09-27 审查回归：并发写入、上传边界、修改范围、配额和版本索引。"""
import tempfile
import threading
import zipfile
from pathlib import Path
from unittest.mock import patch

from base import DBTestCase


class TestRegressions(DBTestCase):
    def test_zip_limits_before_metadata_decompression(self):
        from app import db
        from app.tools import limits
        from app.util import UserError
        path = self.tmp / 'metadata-bomb.docx'
        content_type = next(k for k, v in limits.OOXML_MAIN.items() if v == 'docx')
        with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr('[Content_Types].xml', content_type + ' ' * (2 * 1024 * 1024))
        old = db.get_setting('max_unzipped_mb', 1024)
        try:
            db.set_setting('max_unzipped_mb', 1)
            with patch.object(zipfile.ZipFile, 'open', side_effect=AssertionError('条目不应被解压')):
                with self.assertRaises(UserError) as e:
                    limits.inspect(path, path.name)
                self.assertEqual(e.exception.status, 413)
            # 总大小合法时，类型元数据还有独立上限。
            db.set_setting('max_unzipped_mb', 10)
            with patch.object(zipfile.ZipFile, 'open', side_effect=AssertionError('超大类型条目不应被解压')):
                with self.assertRaises(UserError):
                    limits.sniff(path, path.name)
            with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
                z.writestr('[Content_Types].xml', content_type)
            self.assertEqual(limits.inspect(path, path.name)['kind'], 'docx')
        finally:
            db.set_setting('max_unzipped_mb', old)

    def test_cond_format_count_per_rule(self):
        """同一区域的多条条件格式规则合并在一个 conditionalFormatting 节点里，清点时按规则计数，不能误报缺失。"""
        from openpyxl import Workbook
        from openpyxl.formatting.rule import CellIsRule
        from openpyxl.styles import PatternFill
        from app.tools import ooxml
        wb = Workbook()
        ws = wb.active
        fill = PatternFill(start_color="FF9999", end_color="FF9999", fill_type="solid")
        ws.conditional_formatting.add("D26", CellIsRule(operator="lessThan", formula=["0"], fill=fill))
        ws.conditional_formatting.add("D26", CellIsRule(operator="greaterThan", formula=["0"], fill=fill))
        ws.conditional_formatting.add("E7:E25", CellIsRule(operator="greaterThan", formula=["0"], fill=fill))
        out = self.tmp / "cf.xlsx"
        wb.save(out)
        self.assertEqual(ooxml.inventory(out)["cond_formats"], 3)

    def test_continued_slide_title_follows_language(self):
        from app.pipeline.ppt import _cont
        self.assertEqual(_cont("Market overview"), "Market overview (cont.)")
        self.assertEqual(_cont("市场概况"), "市场概况（续）")

    def test_blob_writes_are_thread_safe(self):
        from app import db, storage
        for kind in ('bytes', 'file', 'move'):
            with self.subTest(kind=kind):
                data = ('concurrent-' + kind).encode()
                gate = threading.Barrier(2)
                replace = storage.os.replace
                results, errors = [], []
                def synchronized_replace(src, dst):
                    gate.wait(timeout=10)
                    return replace(src, dst)
                def write(i):
                    try:
                        if kind == 'bytes':
                            results.append(storage.put_bytes(data))
                        else:
                            src = self.tmp / f'blob-{kind}-{i}'
                            src.write_bytes(data)
                            results.append(storage.put_file(src, move=kind == 'move'))
                    except Exception as exc:
                        errors.append(exc)
                    finally:
                        db.close_thread_conn()
                with patch.object(storage.os, 'replace', synchronized_replace):
                    threads = [threading.Thread(target=write, args=(i,)) for i in range(2)]
                    for t in threads:
                        t.start()
                    for t in threads:
                        t.join(15)
                self.assertFalse(any(t.is_alive() for t in threads))
                self.assertEqual(errors, [])
                self.assertEqual(len(results), 2)
                self.assertEqual(results[0], results[1])
                blob = storage.blob_path(results[0][0])
                self.assertEqual(blob.read_bytes(), data)
                self.assertEqual(blob.stat().st_mode & 0o777, 0o644)
                self.assertEqual(list(blob.parent.glob(blob.name + '.*.tmp')), [])

    def test_ppt_order_and_settings_are_saved(self):
        from app import accounts, jobs, works
        from app.pipeline import edit
        from app.pipeline.context import Ctx
        from app.pipeline.samples import sample_deck
        from app.spec.deck import Deck
        ws = accounts.create_workspace('ppt-regression', 'owner')
        for change in ('move', 'footer', 'aspect', 'title', 'noop'):
            with self.subTest(change=change):
                deck = Deck.model_validate(sample_deck())
                wid = works.create_work(ws, 'ppt', 'Original')
                v = works.add_version(wid, job_id=None, source='test', spec=deck.model_dump(mode='json'))
                ops = {'move': [{'op': 'move_slide', 'slide_id': deck.slides[0].id, 'to': 2}],
                       'footer': [{'op': 'set_deck', 'footer': 'New footer'}],
                       'aspect': [{'op': 'set_deck', 'aspect': '4:3'}],
                       'title': [{'op': 'set_deck', 'title': 'New title'}], 'noop': []}[change]
                job = jobs.enqueue('edit', ws, {'instruction': change, 'base_version_id': v['id']}, work_id=wid)
                ctx = Ctx(job)
                def model(*args, **kw):
                    return kw['validate']({'ops': ops, 'reply': 'ok'})
                def render(ctx, spec, *a, **kw):
                    return spec, {}, []
                with patch.object(ctx.llm, 'json', model), patch.object(edit, 'render_and_fix', render), \
                     patch.object(edit, 'resolve_images'), patch.object(edit, 'LARGE_DECK', 100):
                    result = edit.handle_edit(ctx)
                new = works.current_version(wid)
                if change == 'noop':
                    self.assertTrue(result['no_change'])
                    self.assertEqual(new['id'], v['id'])
                else:
                    self.assertNotEqual(new['id'], v['id'])
                    self.assertEqual(len(new['changed']), len(deck.slides))
                    if change == 'move':
                        self.assertEqual(new['spec']['slides'][2]['id'], deck.slides[0].id)
                    else:
                        self.assertEqual(new['spec'][change], ops[0][change])

    def test_import_notes_cannot_escape_selection(self):
        from pptx import Presentation
        from app import accounts, jobs, storage, works
        from app.pipeline.context import Ctx
        from app.pipeline import importer
        ws = accounts.create_workspace('notes-regression', 'owner')
        path = self.tmp / 'notes.pptx'
        prs = Presentation()
        for i in (1, 2):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = str(i)
            slide.notes_slide.notes_text_frame.text = f'notes {i}'
        prs.save(path)
        sha, _ = storage.put_file(path)
        for scope in ({'type': 'slides', 'ids': [1]}, {'type': 'units', 'ids': ['s1/notes']}):
            wid = works.create_work(ws, 'import_pptx', 'Notes')
            base = works.add_version(wid, job_id=None, source='test', file_sha=sha, manifest={'file_kind': 'pptx'})
            job = jobs.enqueue('import_edit', ws, {'instruction': 'notes', 'scope': scope, 'base_version_id': base['id']}, work_id=wid)
            ctx = Ctx(job)
            def outside(*args, **kw):
                return kw['validate']({'ops': [{'op': 'set_notes', 'index': 2, 'text': 'outside'}]})
            with patch.object(ctx.llm, 'json', outside), patch.object(importer.inplace, 'apply') as apply:
                with self.assertRaises(ValueError):
                    importer.handle_import_edit(ctx)
                apply.assert_not_called()
            def inside(*args, **kw):
                return kw['validate']({'ops': [{'op': 'set_notes', 'index': 1, 'text': 'updated'}]})
            with patch.object(ctx.llm, 'json', inside), patch.object(ctx, 'child', return_value={'pages': [], 'pdf': {}}):
                result = importer.handle_import_edit(ctx)
            v = works.get_version(result['version_id'])
            output = Presentation(storage.blob_path(v['file_sha']))
            self.assertEqual(output.slides[0].notes_slide.notes_text_frame.text, 'updated')
            self.assertEqual(output.slides[1].notes_slide.notes_text_frame.text, 'notes 2')

    def test_stale_and_unsettled_final_jobs_release_quota(self):
        from app import accounts, db, jobs, quota, scheduler
        from app.util import now
        _, tok = accounts.create_token('stale-regression', now() + 3600, quota={'gen_count': 10, 'tokens': 10000})
        ids = []
        for mode in ('failed', 'cancelled', 'retry', 'done'):
            jid = 'regression_' + mode
            with db.tx() as conn:
                quota.reserve(jid, tok['id'], 'gen_doc', 1000, conn=conn)
                jobs.enqueue('gen_doc', tok['workspace_id'], token_id=tok['id'], job_id=jid,
                             max_attempts=2 if mode == 'retry' else 1, conn=conn)
            db.run("UPDATE jobs SET status=?, attempts=1, heartbeat_at=?, cancel_requested=? WHERE id=?",
                   ('done' if mode == 'done' else 'running', now()-300, int(mode == 'cancelled'), jid))
            # 第三个任务保持在重试队列，需要继续保留配额。完成任务模拟 worker 未结算便退出。
            if mode == 'done':
                ids.append(jid)
            elif mode != 'retry':
                ids.append(jid)
        db.insert('usage', {'id': 'regression_usage', 'root_job_id': 'regression_failed',
                            'prompt_tokens': 100, 'completion_tokens': 23, 'created_at': now()})
        scheduler.requeue_stale()
        self.assertEqual(jobs.get('regression_failed')['status'], 'failed')
        self.assertEqual(jobs.get('regression_cancelled')['status'], 'cancelled')
        self.assertEqual(jobs.get('regression_retry')['status'], 'queued')
        for jid in ids:
            self.assertEqual(db.one('SELECT status FROM reservations WHERE job_id=?', (jid,))['status'], 'settled')
        self.assertEqual(db.one('SELECT status FROM reservations WHERE job_id=?', ('regression_retry',))['status'], 'reserved')
        before = quota.usage_summary(tok['id'])
        self.assertEqual(before, {'gen': 2, 'tokens': 1123, 'running': 1})
        scheduler.requeue_stale()
        self.assertEqual(quota.usage_summary(tok['id']), before)
        quota.settle('regression_retry')
        self.assertEqual(db.one('SELECT status FROM reservations WHERE job_id=?', ('regression_retry',))['status'], 'reserved')

    def test_restore_and_copy_preserve_version_search_text(self):
        from app import accounts, db, works
        ws = accounts.create_workspace('search-regression', 'owner')
        wid = works.create_work(ws, 'doc', 'Report')
        old = works.add_version(wid, job_id=None, source='test', spec={}, search_text='original keyword')
        works.add_version(wid, job_id=None, source='test', spec={}, search_text='new keyword')
        restored = works.restore(wid, old['id'])
        self.assertEqual(restored['search_text'], 'original keyword')
        copied = works.duplicate_to(wid, ws)
        for w in (wid, copied):
            self.assertEqual(db.one('SELECT body FROM works_fts WHERE work_id=?', (w,))['body'], 'original keyword')

    def test_legacy_version_search_text_is_rebuilt(self):
        from app import accounts, db, storage, works
        from app.pipeline.samples import sample_document
        from app.spec.document import Document, document_text
        ws = accounts.create_workspace('legacy-search', 'owner')
        doc = Document.model_validate(sample_document())
        wid = works.create_work(ws, 'doc', 'Legacy')
        v = works.add_version(wid, job_id=None, source='test', spec=doc.model_dump(mode='json'))
        db.run('UPDATE versions SET search_text=NULL WHERE id=?', (v['id'],))
        self.assertEqual(works.restore(wid, v['id'])['search_text'], document_text(doc))
        from docx import Document as Word
        path = self.tmp / 'legacy.docx'
        file = Word(); file.add_paragraph('legacy imported keyword'); file.save(path)
        sha, _ = storage.put_file(path)
        wid = works.create_work(ws, 'import_docx', 'Legacy import')
        v = works.add_version(wid, job_id=None, source='test', file_sha=sha)
        db.run('UPDATE versions SET search_text=NULL WHERE id=?', (v['id'],))
        self.assertIn('legacy imported keyword', works.restore(wid, v['id'])['search_text'])
        from openpyxl import Workbook
        path = self.tmp / 'legacy.xlsx'
        book = Workbook(); book.active['A1'] = 'legacy spreadsheet keyword'; book.save(path)
        sha, _ = storage.put_file(path)
        wid = works.create_work(ws, 'import_xlsx', 'Legacy spreadsheet')
        v = works.add_version(wid, job_id=None, source='test', file_sha=sha)
        db.run('UPDATE versions SET search_text=NULL WHERE id=?', (v['id'],))
        self.assertIn('legacy spreadsheet keyword', works.restore(wid, v['id'])['search_text'])

    def test_version_changed_during_render_cannot_be_overwritten(self):
        from app import accounts, works
        from app.util import UserError
        ws = accounts.create_workspace('version-conflict', 'owner')
        wid = works.create_work(ws, 'doc', 'Report')
        old = works.add_version(wid, job_id=None, source='test', search_text='old')
        current = works.add_version(wid, job_id=None, source='test', search_text='current')
        for source in ('chat', 'manual', 'inplace'):
            with self.subTest(source=source), self.assertRaises(UserError) as err:
                works.add_version(wid, job_id=None, source=source, base_version_id=old['id'])
            self.assertEqual(err.exception.status, 409)
        self.assertEqual(works.current_version(wid)['id'], current['id'])
        self.assertEqual(len(works.list_versions(wid)), 2)
        # 主动恢复历史版本仍然允许。
        self.assertEqual(works.restore(wid, old['id'])['search_text'], 'old')

    def test_database_migration_is_safe_when_started_concurrently(self):
        from app import db
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / 'migration.sqlite3'
            conns = [db._connect(path) for _ in range(3)]
            gate = threading.Barrier(len(conns)); errors = []
            def migrate(conn):
                try:
                    gate.wait(timeout=10)
                    db.migrate(conn)
                except Exception as exc:
                    errors.append(exc)
            threads = [threading.Thread(target=migrate, args=(c,)) for c in conns]
            for t in threads: t.start()
            for t in threads: t.join(15)
            try:
                self.assertEqual(errors, [])
                self.assertFalse(any(t.is_alive() for t in threads))
                self.assertEqual(conns[0].execute("SELECT value FROM meta WHERE key='schema_version'").fetchone()[0], str(len(db.MIGRATIONS)))
                self.assertIn('search_text', [r['name'] for r in conns[0].execute('PRAGMA table_info(versions)')])
            finally:
                for c in conns: c.close()

    def test_xlsx_inplace_bad_unit_is_skipped(self):
        from openpyxl import Workbook, load_workbook
        from app.tools import inplace
        src = self.tmp / "u.xlsx"
        wb = Workbook(); wb.active.title = "表"; wb.active["A1"] = 1; wb.save(src)
        dst = self.tmp / "u2.xlsx"
        res = inplace.xlsx_apply(src, dst, [{"op": "set_cell", "unit": "表!A1:B2", "value": 5}, {"op": "set_cell", "unit": "无!A1", "value": 5},
                                           {"op": "set_cell", "unit": "表!A1", "value": 7}])
        self.assertEqual(res["applied"], 1)
        self.assertEqual(len(res["skipped"]), 2)
        self.assertEqual(load_workbook(dst)["表"]["A1"].value, 7)

    def test_xls_plan_errors_are_fed_back(self):
        import pandas as pd
        from app.pipeline.xls import validate_plan
        frames = {"销售": pd.DataFrame({"地区": ["a", "b"], "销售额": [1, 2]})}
        ok = {"sheets": [{"source": "销售", "ops": [{"op": "group", "by": ["地区"], "aggs": [{"column": "销售额", "func": "sum"}]}]}]}
        self.assertIs(validate_plan(frames, ok), ok)
        for bad in ({"sheets": [{"source": "销售", "ops": [{"op": "sort", "by": "不存在的列"}]}]},
                    {"sheets": [{"source": "销售", "ops": [{"op": "explode"}]}]},
                    {"sheets": [{"source": "没有", "ops": []}]}, {"sheets": ["x"]}, {"sheets": [{"source": "销售", "ops": [{"op": "derive", "name": "n", "left": "销售额", "right": "abc"}]}]}):
            with self.subTest(bad=bad), self.assertRaises(ValueError):
                validate_plan(frames, bad)

    def test_blob_rewritten_if_gc_removed_it(self):
        """垃圾回收在"内容已存在"与"登记引用"之间删掉了文件：登记后发现文件不在时重新写入。"""
        from unittest.mock import patch
        from app import db, storage
        data = b"gc-race-content"
        sha, _ = storage.put_bytes(data)
        blob = storage.blob_path(sha)
        self.assertTrue(blob.exists())
        real = db.run
        def run_and_unlink(sql, args=(), conn=None):
            n = real(sql, args, conn)
            if sql.startswith("INSERT OR IGNORE INTO blobs"):
                blob.unlink(missing_ok=True)  # 模拟 gc 恰好此时删掉文件
            return n
        with patch.object(storage.db, "run", run_and_unlink):
            storage.put_bytes(data)
        self.assertEqual(blob.read_bytes(), data)
        src = self.tmp / "gc-src.bin"
        src.write_bytes(data)
        with patch.object(storage.db, "run", run_and_unlink):
            storage.put_file(src, move=True)
        self.assertEqual(blob.read_bytes(), data)
        self.assertFalse(src.exists())

    def test_gc_rechecks_references_before_deleting(self):
        """快照之后新建的引用（例如新上传复用了旧内容）不会被误删。"""
        from app import db, scheduler, storage, works
        ws = "gc-ws"
        data = b"old-content-reused"
        sha, size = storage.put_bytes(data)
        db.run("UPDATE blobs SET created_at=0 WHERE sha=?", (sha,))
        real = scheduler._blob_refs
        state = {"n": 0}
        def refs_then_reference(conn):
            r = real(conn)
            if state["n"] == 0:
                # 快照读完后，工作区才引用这个内容
                storage.create_file_record(ws, sha, size, "reuse.bin", "upload")
            state["n"] += 1
            return r
        from unittest.mock import patch
        with patch.object(scheduler, "_blob_refs", refs_then_reference):
            scheduler.gc_blobs(0)
        self.assertTrue(storage.blob_path(sha).exists())
        self.assertIsNotNone(db.one("SELECT 1 AS x FROM blobs WHERE sha=?", (sha,)))
        db.run("UPDATE files SET deleted_at=1 WHERE sha=?", (sha,))
        scheduler.gc_blobs(0)
        self.assertFalse(storage.blob_path(sha).exists())
        _ = works

    def test_upgrade_preserves_existing_versions(self):
        from app import db
        with tempfile.TemporaryDirectory() as temp:
            conn = db._connect(Path(temp) / 'old.sqlite3')
            try:
                with patch.object(db, 'MIGRATIONS', db.MIGRATIONS[:1]):
                    db.migrate(conn)
                conn.execute("INSERT INTO works(id,workspace_id,kind,title,created_at,updated_at) VALUES('w','ws','doc','old',1,1)")
                conn.execute("INSERT INTO versions(id,work_id,number,source,spec,created_at) VALUES('v','w',1,'generate','{}',1)")
                db.migrate(conn)
                row = dict(conn.execute("SELECT * FROM versions WHERE id='v'").fetchone())
                self.assertEqual(row['spec'], '{}')
                self.assertEqual(row['number'], 1)
                self.assertIsNone(row['search_text'])
                db.migrate(conn)
                self.assertEqual(conn.execute('SELECT COUNT(*) FROM versions').fetchone()[0], 1)
            finally:
                conn.close()
