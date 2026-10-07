"""Docker 镜像冒烟测试：在同一镜像的一次性容器里运行，通过网页接口走一遍“登录 → 上传 → 转换 → 下载”。

    docker compose run --rm -T --no-deps -v "$PWD/tests/smoke_docker.py:/smoke.py:ro" web python /smoke.py http://web:8000

只用镜像里已有的依赖（httpx、python-docx）。管理员账号须已创建为 admin / password1234。
"""
import io
import sys
import time

import httpx

BASE = sys.argv[1] if len(sys.argv) > 1 else "http://web:8000"
TIMEOUT = int(sys.argv[2]) if len(sys.argv) > 2 else 240


def wait_healthy(c: httpx.Client) -> None:
    for _ in range(120):
        try:
            if c.get("/healthz").status_code == 200:
                return
        except httpx.HTTPError:
            pass
        time.sleep(1)
    sys.exit("web 服务 120 秒内没有就绪")


def make_docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_heading("冒烟测试文档", 1)
    for i in range(30):
        d.add_paragraph(f"第 {i + 1} 段：这是用于验证容器内 LibreOffice 转换的中文内容。")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def upload(c: httpx.Client, name: str, data: bytes) -> dict:
    u = c.post("/api/uploads", json={"name": name, "size": len(data)}).raise_for_status().json()
    cs = u["chunk_size"]
    for i in range(u["chunks"]):
        c.put(f"/api/uploads/{u['upload_id']}/{i}", content=data[i * cs:(i + 1) * cs]).raise_for_status()
    return c.post(f"/api/uploads/{u['upload_id']}/complete", json={"purpose": "source"}).raise_for_status().json()


def wait_job(c: httpx.Client, jid: str) -> dict:
    deadline = time.time() + TIMEOUT
    while time.time() < deadline:
        j = c.get(f"/api/jobs/{jid}").raise_for_status().json()
        if j["status"] in ("done", "failed", "cancelled"):
            return j
        time.sleep(2)
    sys.exit(f"任务 {jid} 在 {TIMEOUT} 秒内没有结束（worker 没有认领？）")


def main() -> int:
    c = httpx.Client(base_url=BASE, timeout=30)
    wait_healthy(c)
    c.post("/api/login", json={"username": "admin", "password": "password1234"}).raise_for_status()
    me = c.get("/api/me").raise_for_status().json()
    c.headers["X-CSRF-Token"] = me["csrf"]
    print("已登录，版本", c.get("/api/admin/system").raise_for_status().json().get("version"))

    f = upload(c, "冒烟.docx", make_docx())
    print("已上传", f["name"], f["size"], "字节")
    results = []
    for target in ("pdf", "pdf"):  # 第二次转换应复用常驻 LibreOffice
        j = c.post("/api/jobs", json={"kind": "convert", "params": {"file_id": f["id"], "target": target}}).raise_for_status().json()
        j = wait_job(c, j["id"])
        if j["status"] != "done":
            sys.exit(f"转换任务失败：{j.get('error')}")
        out = j["result"]["files"][0]
        pdf = c.get(f"/api/files/{out['file_id']}/download").raise_for_status().content
        if not pdf.startswith(b"%PDF"):
            sys.exit("下载的结果不是 PDF")
        results.append(len(pdf))
        print(f"转换为 {target} 完成，{len(pdf)} 字节")
    # 导入 + 预览（走 convert 和 preview 两个队列）
    j = c.post("/api/jobs", json={"kind": "import_file", "params": {"file_id": f["id"]}}).raise_for_status().json()
    j = wait_job(c, j["id"])
    if j["status"] != "done":
        sys.exit(f"导入任务失败：{j.get('error')}")
    w = c.get(f"/api/works/{j['result']['work_id']}").raise_for_status().json()
    pages = ((w.get("version") or {}).get("manifest") or {}).get("pages") or []
    print("导入完成，预览页数", len(pages))
    if not pages:
        sys.exit("导入后没有预览页")
    sysinfo = c.get("/api/admin/system").raise_for_status().json()
    print("沙箱：", sysinfo.get("sandbox_effective"), "OCR：", sysinfo.get("ocr_engine"))
    if not sysinfo.get("sandbox_effective"):
        sys.exit("重负载 worker 没有报告沙箱状态")
    print("冒烟测试通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
