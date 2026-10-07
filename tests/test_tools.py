import shutil
import unittest
import zipfile
from pathlib import Path

from base import DBTestCase
from samples import sample_deck, sample_document, sample_workbook


class TestTools(DBTestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        from app.render.docx_render import render_document
        from app.render.pptx_render import render_deck
        from app.render.xlsx_render import render_workbook
        from app.spec.deck import Deck
        from app.spec.document import Document
        from app.spec.workbook import Workbook
        cls.w = cls.tmp / "work"
        cls.w.mkdir()
        cls.docx = cls.w / "doc.docx"
        cls.docres = render_document(Document.model_validate(sample_document()), cls.docx, {}, cls.w / "t")
        cls.pptx = cls.w / "deck.pptx"
        render_deck(Deck.model_validate(sample_deck()), cls.pptx, {}, cls.w / "t")
        cls.xlsx = cls.w / "wb.xlsx"
        render_workbook(Workbook.model_validate(sample_workbook()), cls.xlsx)
        from app.tools import office
        cls.pdf = office.to_pdf(cls.docx, cls.w / "pdf")

    def test_word_native_chart(self):
        """Word 中的图表是原生图表：有图表部件、内嵌可编辑的数据工作簿，包结构自检通过。"""
        import io
        import openpyxl
        from app.tools import ooxml
        z = zipfile.ZipFile(self.docx)
        names = z.namelist()
        self.assertIn("word/charts/chart1.xml", names)
        emb = [n for n in names if n.startswith("word/embeddings/") and n.endswith(".xlsx")]
        self.assertEqual(len(emb), 1)
        wb = openpyxl.load_workbook(io.BytesIO(z.read(emb[0])))
        vals = [c.value for row in wb.active.iter_rows() for c in row]
        self.assertIn(2.51, vals)
        self.assertIn("c:chart", z.read("word/document.xml").decode())
        self.assertEqual(ooxml.validate_package(self.docx), [])
        self.assertEqual(ooxml.inventory(self.docx).get("charts"), 1)

    def test_sniff_and_limits(self):
        from app.tools import limits
        from app.util import UserError
        self.assertEqual(limits.sniff(self.docx, "a.docx"), "docx")
        self.assertEqual(limits.sniff(self.pptx, "x.bin"), "pptx")
        self.assertEqual(limits.sniff(self.xlsx, "y"), "xlsx")
        self.assertEqual(limits.inspect(self.pdf, "a.pdf")["kind"], "pdf")
        fake = self.w / "fake.docx"
        fake.write_bytes(b"MZ\x90\x00 not a document")
        with self.assertRaises(UserError):
            limits.sniff(fake, "fake.docx")
        # ZIP 炸弹：内容类型合法，但解压后超过上限
        bomb = self.w / "bomb.docx"
        with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("[Content_Types].xml", '<Types><Override ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
            z.writestr("word/big.bin", b"\0" * (1100 * 1024 * 1024))
        with self.assertRaises(UserError):
            limits.inspect(bomb, "bomb.docx")
        bomb.unlink()
        from PIL import Image
        big = self.w / "big.png"
        Image.new("L", (12000, 9000)).save(big)
        with self.assertRaises(UserError):
            limits.inspect(big, "big.png")

    def test_pdf_tools(self):
        from app.tools import pdf
        n = pdf.page_count(self.pdf)
        self.assertGreaterEqual(n, 2)
        m = pdf.merge([self.pdf, self.pdf], self.w / "merged.pdf")
        self.assertEqual(pdf.page_count(m), 2 * n)
        parts = pdf.split(m, self.w / "split", ranges=f"1-{n},{n + 1}-")
        self.assertEqual([pdf.page_count(p) for p in parts], [n, n])
        every = pdf.split(m, self.w / "split2", mode="every", every=1)
        self.assertEqual(len(every), 2 * n)
        r = pdf.rotate(self.pdf, self.w / "rot.pdf", "1", 90)
        import pikepdf
        with pikepdf.open(r) as p:
            self.assertEqual(int(p.pages[0].obj.get("/Rotate", 0)), 90)
        ro = pdf.reorder(m, self.w / "ro.pdf", [2, 1])
        self.assertEqual(pdf.page_count(ro), 2)
        rep = pdf.compress(m, self.w / "small.pdf", "ebook")
        self.assertLessEqual(rep["after"], rep["before"])
        # Word 中的图表现在是原生矢量图表；用演示文稿的 PDF（含瀑布图等图片图表）测试图片提取
        from app.tools import office
        deck_pdf = office.to_pdf(self.pptx, self.w / "deckpdf")
        imgs = pdf.extract_images(deck_pdf, self.w / "imgs")
        self.assertGreaterEqual(len(imgs), 1, "图片图表应被提取")
        text, _ = pdf.extract_text(self.pdf)
        self.assertIn("尿酸", text)
        dests = pdf.named_dest_pages(self.pdf)
        heading_ids = [h["id"] for h in self.docres.headings]
        self.assertTrue(all(f"dw{h}" in dests for h in heading_ids), "每个标题都应有命名目标")

    def test_conversions(self):
        from app.tools import convert
        out = self.w / "conv"
        outs, rep = convert.convert(self.pdf, "pdf", "docx", out / "a")
        self.assertTrue(outs[0].exists())
        from docx import Document
        txt = "\n".join(p.text for p in Document(outs[0]).paragraphs)
        self.assertIn("尿酸", txt)
        outs, rep = convert.convert(self.pdf, "pdf", "xlsx", out / "b")
        self.assertGreaterEqual(rep["tables"], 1)
        outs, _ = convert.convert(self.pdf, "pdf", "pptx_image", out / "c")
        from pptx import Presentation
        self.assertGreaterEqual(len(Presentation(outs[0]).slides), 2)
        outs, _ = convert.convert(self.pdf, "pdf", "png", out / "d")
        self.assertTrue(outs[0].name.endswith(".zip"))
        outs, _ = convert.convert(self.pptx, "pptx", "pdf", out / "e")
        self.assertTrue(outs[0].exists())
        md = self.w / "note.md"
        md.write_text("# 标题\n\n正文 **加粗**\n\n| a | b |\n|---|---|\n| 1 | 2 |\n", encoding="utf-8")
        outs, _ = convert.convert(md, "md", "docx", out / "f")
        self.assertIn("标题", "\n".join(p.text for p in Document(outs[0]).paragraphs))
        outs, _ = convert.convert(md, "md", "pdf", out / "g")
        self.assertTrue(outs[0].exists())
        outs, _ = convert.convert(self.docx, "docx", "md", out / "h")
        self.assertIn("尿酸", outs[0].read_text(encoding="utf-8"))
        csv = self.w / "d.csv"
        csv.write_text("名称,数量\n甲,1\n乙,2\n", encoding="gbk")
        outs, _ = convert.convert(csv, "csv", "xlsx", out / "i")
        outs2, _ = convert.convert(outs[0], "xlsx", "csv", out / "j")
        self.assertIn("甲", outs2[0].read_text(encoding="utf-8-sig"))
        from PIL import Image
        png = self.w / "a.png"
        Image.new("RGBA", (300, 200), (255, 0, 0, 128)).save(png)
        outs, _ = convert.convert(png, "png", "pdf", out / "k")
        self.assertTrue(outs[0].exists())
        self.assertTrue(any(t["target"] == "pptx_rebuild" for t in convert.targets_for("pdf")))

    def test_macro_strip_and_inventory(self):
        from app.tools import limits, ooxml
        # 构造一个带宏的 pptm：在 pptx 基础上加入 vbaProject.bin 与关系
        pptm = self.w / "m.pptm"
        with zipfile.ZipFile(self.pptx) as zin, zipfile.ZipFile(pptm, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename == "[Content_Types].xml":
                    data = data.decode().replace("presentationml.presentation.main+xml", "PLACEHOLDER").replace(
                        "application/vnd.openxmlformats-officedocument.PLACEHOLDER", "application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml")
                    data = data.replace("</Types>", '<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>').encode()
                if item.filename == "ppt/_rels/presentation.xml.rels":
                    data = data.decode().replace("</Relationships>", '<Relationship Id="rIdVba" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>').encode()
                zout.writestr(item, data)
            zout.writestr("ppt/vbaProject.bin", b"\x00fake-vba")
        self.assertEqual(limits.sniff(pptm, "m.pptm"), "pptm")
        clean = ooxml.strip_macros(pptm, self.w / "clean.pptx")
        self.assertEqual(limits.sniff(clean, "clean.pptx"), "pptx")
        with zipfile.ZipFile(clean) as z:
            self.assertNotIn("ppt/vbaProject.bin", z.namelist())
        self.assertEqual(ooxml.validate_package(clean), [])
        from pptx import Presentation
        Presentation(clean)
        inv = ooxml.inventory(self.xlsx)
        self.assertEqual(inv.get("charts"), 1)
        self.assertEqual(ooxml.compare_inventory({"charts": 2}, {"charts": 1}), ["图表从 2 个变为 1 个"])

    def test_resident_office(self):
        """常驻 LibreOffice：复用同一实例；坏文件仍报准确错误且不影响实例；实例被结束后自动重启；超时会结束实例。"""
        import os
        import signal
        from app.tools import lo_resident, office, sandbox
        if not lo_resident.enabled():
            self.skipTest("系统 Python 无法导入 uno，或当前沙箱不是 rlimit")
        out = self.w / "resident"
        a = office.to_pdf(self.docx, out / "a")
        inst = lo_resident._instance()
        pid = inst.proc.pid
        b = office.to_pdf(self.pptx, out / "b")
        self.assertEqual(inst.proc.pid, pid, "第二次转换应复用同一个常驻实例")
        self.assertGreater(a.stat().st_size, 1000)
        self.assertGreater(b.stat().st_size, 1000)
        x = office.recalc_copy(self.xlsx, out / "c")
        self.assertEqual(x.suffix, ".xlsx")
        bad = self.w / "bad.docx"
        bad.write_bytes(b"PK\x03\x04 broken")
        with self.assertRaises(sandbox.ToolError):
            office.to_pdf(bad, out / "d")
        self.assertTrue(inst.alive(), "坏文件不应结束常驻实例")
        self.assertTrue(lo_resident.enabled())
        os.killpg(inst.proc.pid, signal.SIGKILL)
        inst.proc.wait()
        c = office.to_pdf(self.docx, out / "e")
        self.assertTrue(c.exists())
        self.assertNotEqual(inst.proc.pid, pid, "实例被结束后应重新启动")
        from unittest import mock
        with mock.patch.object(lo_resident.sandbox, "run", side_effect=sandbox.ToolError("运行超时")):
            with self.assertRaises(sandbox.ToolError):
                office.to_pdf(self.docx, out / "f")
        self.assertFalse(inst.alive(), "超时后应结束常驻实例")

    def test_sandbox_timeout_and_cancel(self):
        import time
        from app.tools import sandbox
        with self.assertRaises(sandbox.ToolError):
            sandbox.run(["sleep", "5"], self.w, timeout=1)
        t0 = time.time()
        with self.assertRaises(sandbox.ToolCancelled):
            sandbox.run(["sleep", "5"], self.w, timeout=30, cancel=lambda: time.time() - t0 > 0.6)
        self.assertLess(time.time() - t0, 3)


if __name__ == "__main__":
    unittest.main()
